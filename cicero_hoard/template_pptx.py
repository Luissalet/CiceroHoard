"""Native editable export using source-slide templates and their OOXML graph."""
from __future__ import annotations

import io
import hashlib
from copy import deepcopy
from dataclasses import replace
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt
from pptx.dml.color import MSO_COLOR_TYPE

from .errors import CiceroError
from .layouts import build_plan, Ctx, Para, fit_size
from .templates import PX, TemplateBinding, _presentation, is_slide_number
from .themes import theme_of

RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _clone(prs, source, excluded):
    slide=prs.slides.add_slide(source.slide_layout)
    for key, value in source._element.attrib.items():
        slide._element.set(key, value)
    override = source._element.find(qn('p:clrMapOvr'))
    if override is not None:
        old = slide._element.find(qn('p:clrMapOvr'))
        if old is not None:slide._element.remove(old)
        slide._element.append(deepcopy(override))
    for s in list(slide.shapes):s._element.getparent().remove(s._element)
    elements=[deepcopy(s._element) for s in source.shapes if s.shape_id not in excluded]
    bg=source._element.cSld.find(qn('p:bg'))
    if bg is not None:
        bg=deepcopy(bg);slide._element.cSld.insert(0,bg);elements.append(bg)
    required={v for el in elements for node in el.iter() for k,v in node.attrib.items() if k.startswith('{'+RNS+'}')}
    mapping={}
    for rid in required:
        rel=source.part.rels[rid]
        mapping[rid]=slide.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,is_external=rel.is_external)
    for el in elements:
        for node in el.iter():
            for key,value in list(node.attrib.items()):
                if key.startswith('{'+RNS+'}') and value in mapping:node.set(key,mapping[value])
        if el.tag!=qn('p:bg'):slide.shapes._spTree.insert_element_before(el,'p:extLst')
    return slide


def _template_theme(source, fallback):
    theme=deepcopy(fallback)
    rel=next((r for r in source.slide_layout.slide_master.part.rels.values() if r.reltype==RT.THEME),None)
    if rel:
        root=etree.fromstring(rel.target_part.blob)
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        colors={}
        for node in root.findall('.//a:clrScheme/*',ns):
            color=node[0] if len(node) else None
            if color is not None:colors[etree.QName(node).localname]='#'+(color.get('val') if color.tag==qn('a:srgbClr') else color.get('lastClr','000000'))
        for key,slot in [('background','lt1'),('surface','lt2'),('text','dk2'),('muted','dk2'),('accent','accent1'),('accent2','accent5')]:
            if slot in colors:theme['colors'][key]=colors[slot]
        for key,slot in [('heading','majorFont'),('body','minorFont')]:
            font=root.find(f'.//a:fontScheme/a:{slot}/a:latin',ns)
            if font is not None and font.get('typeface'):theme['fonts'][key]=font.get('typeface')
    return theme


def _replace_text(shape, value, theme, prefix=''):
    tf=shape.text_frame;paragraph=tf.paragraphs[0]
    runs=[r for p in tf.paragraphs for r in p.runs if r.text.strip()]
    first=deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    last=deepcopy(runs[-1]._r.rPr) if runs and runs[-1]._r.rPr is not None else first
    ppr=deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    base=next((r.font.size.pt/.75 for r in reversed(runs) if r.font.size),48)
    texts=([prefix] if prefix else [])+[value]
    size,_=fit_size(lambda n:[Para(t,n) for t in texts],max(1,shape.width/PX-20),max(1,shape.height/PX-10),round(base),18,Ctx(theme))
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
    for i,text in enumerate(texts):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if p._p.pPr is not None:p._p.remove(p._p.pPr)
        if ppr is not None:p._p.insert(0,deepcopy(ppr))
        r=p.add_run();r.text=text;style=first if prefix and i==0 else last
        if style is not None:r._r.insert(0,deepcopy(style))
        r.font.size=Pt(size*.75)


def build_template_pptx(deck, *, assets):
    from .export_pptx import _add_box
    descriptor=deck['template'];path=Path(descriptor['path'])
    try:data=path.read_bytes()
    except OSError:raise CiceroError('The attached template snapshot is missing; attach it again.') from None
    if hashlib.sha256(data).hexdigest()!=descriptor['sha256']:raise CiceroError('The attached template snapshot changed; attach it again.')
    prs=_presentation(data);sources=list(prs.slides)
    for n,item in enumerate(deck['slides'],1):
        binding=TemplateBinding.model_validate(descriptor['bindings'][item['layout']]);source=sources[binding.source_slide-1]
        theme=_template_theme(source,theme_of(deck));keep=set(binding.keep_text_shapes)
        title=next((s for s in source.shapes if s.shape_id==binding.title_shape),None)
        subtitle=next((s for s in source.shapes if s.shape_id==binding.subtitle_shape),None)
        numbers={s.shape_id for s in source.shapes if is_slide_number(s)}
        slots=[s for s in source.shapes if ((s.has_text_frame and (s.text.strip() or s.is_placeholder)) or s.has_chart or s.has_table)
               and s.shape_id not in keep|numbers|{binding.title_shape,binding.subtitle_shape}]
        text_slots=[s for s in slots if s.has_text_frame and s.text.strip()]
        if text_slots:
            body=max(text_slots,key=lambda s:len(s.text))
            runs=[r for p in body.text_frame.paragraphs for r in p.runs if r.text.strip()]
            if runs:
                if runs[0].font.name:theme['fonts']['body']=runs[0].font.name
                if runs[0].font.color.type==MSO_COLOR_TYPE.RGB:theme['colors']['text']='#'+str(runs[0].font.color.rgb)
        excluded=set(binding.remove_shapes)|{s.shape_id for s in slots}
        slide=_clone(prs,source,excluded)
        cloned={s.shape_id:s for s in slide.shapes}
        if title is not None and title.shape_id in cloned:_replace_text(cloned[title.shape_id],item['title'],theme,binding.prefix)
        if subtitle is not None and subtitle.shape_id in cloned:_replace_text(cloned[subtitle.shape_id],item.get('subtitle') or '',theme)
        for sid in numbers:
            if sid in cloned:
                s=cloned[sid];s.width=min(s.width,prs.slide_width-s.left);_replace_text(s,f'{n:02d}',theme)
        if binding.content_frame:frame=binding.content_frame
        elif slots:
            x=min(s.left for s in slots)/PX;y=min(s.top for s in slots)/PX
            frame=[x,y,max(s.left+s.width for s in slots)/PX-x,max(s.top+s.height for s in slots)/PX-y]
        else:
            y=((title.top+title.height)/PX+25) if title is not None else prs.slide_height/PX*.22
            frame=[prs.slide_width/PX*.06,y,prs.slide_width/PX*.88,max(1,prs.slide_height/PX*.85-y)]
        plan=build_plan(item,deck_title=deck['title'],theme=theme,number=n,total=len(deck['slides']),lang=deck.get('language','es'),assets=assets)
        boxes=[b for b in plan.boxes if b.kind!='rect' and b.role!='footer' and (b.role!='title' or title is None)]
        # Subtitle lives in its own template slot; never duplicate it as body.
        if subtitle is not None:boxes=[b for b in boxes if [p.text for p in b.paras]!=[item.get('subtitle') or '']]
        if boxes:
            bx=min(b.x for b in boxes);by=min(b.y for b in boxes);bw=max(b.x+b.w for b in boxes)-bx;bh=max(b.y+b.h for b in boxes)-by
            fx,fy,fw,fh=frame;sx=fw/max(1,bw);sy=fh/max(1,bh);scale=min(sx,sy)
            for box in boxes:
                moved=replace(box,x=fx+(box.x-bx)*sx,y=fy+(box.y-by)*sy,w=box.w*sx,h=box.h*sy,
                              paras=[replace(p,size=p.size*scale,indent=round(p.indent*scale),space_after=p.space_after*scale) for p in box.paras])
                _add_box(slide,moved,{**theme, 'chart_font_scale': scale},deck.get('language','es'),assets,None)
        slide.notes_slide.notes_text_frame.text=item.get('notes') or ''
    # Source slides were examples, not additional pages in the authored deck.
    for source in sources:
        entry=next(e for e in prs.slides._sldIdLst if prs.part.related_part(e.rId) is source.part)
        prs.part.drop_rel(entry.rId);prs.slides._sldIdLst.remove(entry)
    prs.core_properties.title=deck['title'];prs.core_properties.author="Cicero's Hoard"
    out=io.BytesIO();prs.save(out);return out.getvalue()
