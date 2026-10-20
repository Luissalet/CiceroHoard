"""Editable native evidence: specialist charts, graph objects and tables."""
from __future__ import annotations
from pptx.chart.data import CategoryChartData,XyChartData,BubbleChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_MARKER_STYLE,XL_TICK_LABEL_POSITION,XL_LEGEND_POSITION,XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
from .layouts import fmt_num,Box,Para
from .themes import hex_to_rgb,contrast


def add_table(slide,box,theme,lang):
    from .export_pptx import _emu,_rgb,_fill_text
    data=box.data;rows=[data['headers'],*data['rows']]
    frame=slide.shapes.add_table(len(rows),len(rows[0]),_emu(box.x),_emu(box.y),_emu(box.w),_emu(box.h))
    size=max(13,min((data.get('font_size') or 24)*theme.get('chart_font_scale',1),box.h/len(rows)*.46))
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            cell=frame.table.cell(i,j);cell.margin_left=cell.margin_right=_emu(10);cell.margin_top=cell.margin_bottom=_emu(6)
            cell.fill.solid();cell.fill.fore_color.rgb=_rgb(theme['colors']['accent'] if i==0 else theme['colors']['surface'] if i-1 in (data.get('highlight_rows') or []) else theme['colors']['background'])
            _fill_text(cell.text_frame,Box('text',0,0,1,1,paras=[Para(str(value),size,bold=i==0,color='text')]),theme,lang)
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.color.rgb=_rgb('#FFFFFF' if i==0 else theme['colors']['text'])
    return frame


def add_graph_object(slide,box,theme):
    from .export_pptx import _emu,_rgb
    if box.kind=='node':
        shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,_emu(box.x),_emu(box.y),_emu(box.w),_emu(box.h))
        shape.adjustments[0]=min(.08,6/max(1,min(box.w,box.h)))
        shape.fill.solid();shape.fill.fore_color.rgb=_rgb(theme['colors'][box.fill or 'background'])
        shape.line.color.rgb=_rgb(theme['colors']['accent']);shape.line.width=Pt(1.1)
        return shape
    cw,ch=box.data['canvas'];points=box.data['points'];shapes=[]
    for i,(a,z) in enumerate(zip(points,points[1:])):
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,_emu(box.x+a['x']*box.w/cw),_emu(box.y+a['y']*box.h/ch),
            _emu(box.x+z['x']*box.w/cw),_emu(box.y+z['y']*box.h/ch))
        line.line.color.rgb=_rgb(theme['colors']['accent']);line.line.width=Pt(1.5)
        if i==len(points)-2:
            end=etree.SubElement(line.line._get_or_add_ln(),qn('a:tailEnd'));end.set('type','triangle');end.set('w','sm');end.set('len','sm')
        shapes.append(line)
    return shapes


def add_chart(slide,box,theme,lang):
    from .export_pptx import _emu,_rgb,_add_box,axis_number_format
    b=box.chart;kind=b['chart'];series=b['series'];cats=b['categories'];palette=[theme['colors']['accent'],theme['colors']['accent2'],'#667B89','#C16B36']
    if not cats or not series:
        from .errors import CiceroError
        raise CiceroError('Add chart categories and series before exporting editable evidence.')
    captions=[]
    if kind in ('pie','donut','funnel'):
        captions.extend(v for v in (b.get('y_label'),b.get('unit')) if v)
    elif b.get('y_label') and b.get('unit') and b['y_label']!=b['unit']:captions.append(b['unit'])
    if b.get('x_label'):captions.append(b['x_label'])
    context=[f"{cats[n['category']]}: {n['text']}" for n in b.get('annotations',[])]
    context += [f"{r['label']}: {fmt_num(r['value'],lang)}" for r in b.get('references',[])]
    if context:captions.append(' | '.join(context))
    plot_height=max(100,box.h-30*len(captions))
    for i,caption in enumerate(captions):
        _add_box(slide,Box('text',box.x,box.y+plot_height+i*30,box.w,28,
            paras=[Para(caption,18,color='muted',align='center')]),theme,lang,None,None)
    if kind=='heatmap':
        values=[v for s in series for v in s['values'] if v is not None] or [0];lo=min(values);hi=max(values)
        data={'headers':['',*[s['name'] for s in series]],'rows':[[cat,*[fmt_num(s['values'][j],lang) if s['values'][j] is not None else '' for s in series]] for j,cat in enumerate(cats)]}
        frame=add_table(slide,Box('table',box.x,box.y,box.w,plot_height,data=data),theme,lang)
        start=hex_to_rgb('#EDF4F7');end=hex_to_rgb(theme['colors']['accent'])
        for i in range(len(cats)):
            for j,s in enumerate(series):
                value=s['values'][i]
                if value is None:continue
                t=(value-lo)/(hi-lo or 1);color='#'+''.join(f'{round(a+(z-a)*t):02X}' for a,z in zip(start,end))
                cell=frame.table.cell(i+1,j+1);cell.fill.fore_color.rgb=_rgb(color)
                cell.text_frame.paragraphs[0].runs[0].font.color.rgb=_rgb('#FFFFFF' if contrast('#FFFFFF',color)>=4.5 else '#172F3B')
        return frame
    if kind=='funnel':
        pairs=[(cat,v) for cat,v in zip(cats,series[0]['values']) if v is not None]
        maximum=max([v for _,v in pairs] or [1]);height=plot_height/max(1,len(pairs))-10;shapes=[]
        for i,(cat,value) in enumerate(pairs):
            width=box.w*.62*value/(maximum or 1);left=box.x+(box.w*.62-width)/2;top=box.y+i*(height+10)
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,_emu(left),_emu(top),_emu(width),_emu(height))
            shape.fill.solid();shape.fill.fore_color.rgb=_rgb(palette[i%len(palette)]);shape.line.fill.background();shapes.append(shape)
            shapes.append(_add_box(slide,Box('text',box.x+box.w*.68,top,box.w*.3,height,anchor='middle',paras=[Para(f'{cat}\n{fmt_num(value,lang)}',24,bold=True)]),theme,lang,None,None))
        return shapes
    xy=kind in ('scatter','bubble','dot','dumbbell','band') or (kind in ('line','area') and any(s.get('x_values') for s in series))
    if xy:
        data=BubbleChartData() if kind=='bubble' else XyChartData()
        if kind=='dumbbell':
            for i,cat in enumerate(cats):
                s=data.add_series(cat)
                for model in series:
                    if model['values'][i] is not None:s.add_data_point(model['values'][i],len(cats)-i)
        else:
            for model in series:
                combinations=[(model['name'],model['values'])]
                if kind=='band':combinations+=[(model['name']+' lower',model['lower']),(model['name']+' upper',model['upper'])]
                for name,values in combinations:
                    s=data.add_series(name)
                    for i,value in enumerate(values):
                        if value is None:
                            if kind in ('band','line','area') and i+1<len(values):s=data.add_series(name+' gap '+str(i))
                            continue
                        x=model.get('x_values',[i+1 for i in range(len(values))])[i]
                        if kind=='dot':x=value;value=len(cats)-i
                        if kind=='bubble':s.add_data_point(x,value,(model.get('sizes') or [10]*len(values))[i])
                        else:s.add_data_point(x,value)
        ctype=XL_CHART_TYPE.BUBBLE if kind=='bubble' else XL_CHART_TYPE.XY_SCATTER if kind in ('scatter','dot') else XL_CHART_TYPE.XY_SCATTER_LINES
    else:
        data=CategoryChartData();data.categories=cats
        for model in series:data.add_series(model['name'],model['values'])
        ctype={'bar_horizontal':XL_CHART_TYPE.BAR_CLUSTERED,'area':XL_CHART_TYPE.AREA,'donut':XL_CHART_TYPE.DOUGHNUT,
            'line':XL_CHART_TYPE.LINE_MARKERS,'pie':XL_CHART_TYPE.PIE}.get(kind,XL_CHART_TYPE.COLUMN_CLUSTERED)
        if b.get('stacked'):
            ctype={'bar':XL_CHART_TYPE.COLUMN_STACKED,'bar_horizontal':XL_CHART_TYPE.BAR_STACKED,
                'area':XL_CHART_TYPE.AREA_STACKED,'line':XL_CHART_TYPE.LINE_STACKED}.get(kind,ctype)
        if kind=='waterfall':
            data=CategoryChartData();data.categories=cats;running=0;parts=[[] for _ in range(6)]
            for value in series[0]['values']:
                value=value or 0;nxt=running+value;lo=min(running,nxt);hi=max(running,nxt)
                pos=max(0,hi)-max(0,lo);neg=min(0,lo)-min(0,hi)
                for dest,v in zip(parts,[max(0,lo),min(0,hi),pos if value>=0 else None,neg if value>=0 else None,pos if value<0 else None,neg if value<0 else None]):dest.append(v)
                running=nxt
            for name,values in zip(['Positive offset','Negative offset','Increase','Increase below zero','Decrease','Decrease below zero'],parts):data.add_series(name,values)
            ctype=XL_CHART_TYPE.COLUMN_STACKED
    frame=slide.shapes.add_chart(ctype,_emu(box.x),_emu(box.y),_emu(box.w),_emu(plot_height),data)
    chart=frame.chart;ids=list(dict.fromkeys(x.get('val') for x in chart._chartSpace.iter(qn('c:axId'))));mapping={v:str(100+i) for i,v in enumerate(ids)}
    for el in chart._chartSpace.iter():
        if el.tag in (qn('c:axId'),qn('c:crossAx')):el.set('val',mapping[el.get('val')])
    points=b['font_size']*.75 if b.get('font_size') else theme.get('chart_font_points',18);chart.font.size=Pt(points);chart.font.name=theme['fonts']['body']
    chart.font.color.rgb=_rgb(theme['colors']['text']);chart.has_title=False
    show_legend=b.get('show_legend')
    chart.has_legend=(len(series)>1 if show_legend is None else show_legend) and kind not in ('dumbbell','dot','funnel','waterfall')
    if chart.has_legend:
        chart.legend.font.size=Pt(points);chart.legend.position=XL_LEGEND_POSITION.TOP;chart.legend.include_in_layout=False
    for i,s in enumerate(chart.series):
        color=series[min(i,len(series)-1)].get('color') or palette[i%len(palette)]
        s.format.fill.solid();s.format.fill.fore_color.rgb=_rgb(color);s.format.line.color.rgb=_rgb(color)
        if xy and kind!='bubble':
            s.marker.style=XL_MARKER_STYLE.CIRCLE;s.marker.size=9
            s.marker.format.fill.solid();s.marker.format.fill.fore_color.rgb=_rgb(color)
            s.marker.format.line.color.rgb=_rgb(color)
            if kind=='dumbbell':
                s.format.line.color.rgb=_rgb('#A7B9C4');s.format.line.width=Pt(2)
                for j,point in enumerate(s.points):
                    ink=series[j].get('color') or palette[j%2]
                    point.marker.format.fill.solid();point.marker.format.fill.fore_color.rgb=_rgb(ink);point.marker.format.line.color.rgb=_rgb(ink)
        if kind=='waterfall':
            if i<2:s.format.fill.background();s.format.line.fill.background()
            else:
                s.format.fill.fore_color.rgb=_rgb(palette[0 if i<4 else 1]);s.invert_if_negative=False
        if kind in ('scatter','dot'):s.format.line.fill.background()
        if kind=='bar_horizontal' and b.get('highlight'):
            for j,point in enumerate(s.points):
                point.format.fill.solid();point.format.fill.fore_color.rgb=_rgb(color if j in b['highlight'] else '#CFD9E0')
    observed=[v for s in series for v in s['values'] if v is not None]
    number_format='#,##0' if all(float(v).is_integer() for v in observed) else '#,##0.###'
    for axis_name in ('category_axis','value_axis'):
        try:
            axis=getattr(chart,axis_name);axis.tick_labels.font.size=Pt(points);axis.tick_labels.font.color.rgb=_rgb(theme['colors']['text']);axis.tick_labels.font.name=theme['fonts']['body']
            axis_values=[x for s in series for x in (s.get('x_values') or [])] if axis_name=='category_axis' else observed
            axis.tick_labels.number_format=axis_number_format(axis_values,from_zero=kind not in ('dot','dumbbell'));axis.tick_labels.number_format_is_linked=False
            axis.major_tick_mark=XL_TICK_MARK.NONE;axis.minor_tick_mark=XL_TICK_MARK.NONE
            if axis.has_major_gridlines:
                axis.major_gridlines.format.line.color.rgb=_rgb('#DFE6EB');axis.major_gridlines.format.line.width=Pt(.5)
            if axis_name=='value_axis' and b.get('zero_baseline',True) and kind not in ('dot','dumbbell'):
                values=[v for s in series for v in s['values'] if v is not None]
                if values and min(values)>=0:axis.minimum_scale=0
            if kind=='bar_horizontal' and axis_name=='category_axis':axis.reverse_order=True
            # Negative values put the axis crossing inside the plot: keep both axes' labels at the plot edge, clear of the data.
            if kind not in ('dot','dumbbell','pie','donut') and observed and min(observed)<0:
                axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
            if kind in ('dot','dumbbell') and axis_name=='value_axis':
                axis.minimum_scale=.5;axis.maximum_scale=len(cats)+.5;axis.tick_label_position=XL_TICK_LABEL_POSITION.NONE
                axis.has_major_gridlines=False
        except (ValueError,AttributeError):pass
    if kind in ('pie','donut'):
        chart.plots[0].vary_by_categories=True
        for j,point in enumerate(chart.series[0].points):
            point.format.fill.solid();point.format.fill.fore_color.rgb=_rgb(palette[j%len(palette)])
    show_values=b.get('show_values')
    if (len(cats)*len(series)<=14 if show_values is None else show_values) and not xy and kind not in ('waterfall','area'):
        from pptx.enum.chart import XL_DATA_LABEL_POSITION
        chart.plots[0].has_data_labels=True
        labels=chart.plots[0].data_labels;labels.number_format=number_format;labels.number_format_is_linked=False;labels.font.size=Pt(points);labels.font.color.rgb=_rgb(theme['colors']['text'])
        labels.show_value=True;labels.show_category_name=kind in ('pie','donut')
        labels.position=(XL_DATA_LABEL_POSITION.CENTER if b.get('stacked') and kind in ('bar','bar_horizontal') else XL_DATA_LABEL_POSITION.OUTSIDE_END if kind in ('bar','bar_horizontal') else
            XL_DATA_LABEL_POSITION.ABOVE if kind=='line' else XL_DATA_LABEL_POSITION.BEST_FIT if kind=='pie' else None)
        labels.font.name=theme['fonts']['body']
    if kind in ('dot','dumbbell'):
        observed=[v for s in series for v in s['values'] if v is not None] or [0]
        maximum=max(0,max(observed))*1.15;minimum=min(0,min(observed))
        if b.get('zero_baseline') is False:
            low,high=min(observed),max(observed);padding=(high-low or abs(high) or 1)*.15
            minimum=low-padding;maximum=high+padding
        if maximum==minimum:maximum=minimum+1
        chart.category_axis.minimum_scale=minimum;chart.category_axis.maximum_scale=maximum
        plot=chart._chartSpace.find('.//'+qn('c:plotArea'));layout=plot.find(qn('c:layout'))
        if layout is None:layout=etree.SubElement(plot,qn('c:layout'))
        manual=etree.SubElement(layout,qn('c:manualLayout'))
        for key,value in [('layoutTarget','inner'),('xMode','factor'),('yMode','factor'),('wMode','factor'),('hMode','factor'),('x','.20'),('y','.12'),('w','.75'),('h','.76')]:
            el=etree.SubElement(manual,qn('c:'+key));el.set('val',value)
        for i,cat in enumerate(cats):
            _add_box(slide,Box('text',box.x,box.y+plot_height*(.12+.76*(i+.5)/len(cats)),box.w*.18,40,paras=[Para(cat,23)]),theme,lang,None,None)
        if kind=='dumbbell':
            for i,s in enumerate(series):
                ink=s.get('color') or palette[i];legend_theme={**theme,'colors':{**theme['colors'],'legend_ink':ink}}
                _add_box(slide,Box('text',box.x+180+i*260,box.y-32,250,32,paras=[Para(s['name'],24,bold=True,color='legend_ink')]),legend_theme,lang,None,None)
            if b.get('show_values',True):
                from pptx.enum.chart import XL_DATA_LABEL_POSITION
                from pptx.chart.datalabel import DataLabels
                for j,s in enumerate(chart.series):
                    labels=DataLabels(s._element.get_or_add_dLbls())
                    labels.show_category_name=False;labels.show_value=False;labels.show_series_name=False
                    labels.font.size=Pt(points);labels.font.name=theme['fonts']['body'];labels.font.color.rgb=_rgb(theme['colors']['text'])
                    # Each point carries its own text: readers disagree on what "category name" means in XY charts
                    # (one prints the X value, another the series name), and the deck language sets the decimal mark.
                    row=[model['values'][j] for model in series if model['values'][j] is not None]
                    for i,point in enumerate(s.points):
                        point.data_label.position=XL_DATA_LABEL_POSITION.ABOVE if i==0 else XL_DATA_LABEL_POSITION.BELOW
                        frame=point.data_label.text_frame;frame.text=fmt_num(row[i],lang)
                        run=frame.paragraphs[0].runs[0];run.font.size=Pt(points);run.font.name=theme['fonts']['body'];run.font.color.rgb=_rgb(theme['colors']['text'])
    unit=b.get('y_label') or b.get('unit')
    if unit and kind not in ('pie','donut','funnel'):
        axis=chart.value_axis;axis.has_title=True
        tf=axis.axis_title.text_frame;tf.text=unit
        for run in tf.paragraphs[0].runs:
            run.font.name=theme['fonts']['body'];run.font.size=Pt(points);run.font.color.rgb=_rgb(theme['colors']['text'])
    return frame
