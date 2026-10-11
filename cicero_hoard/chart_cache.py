"""Reconcile native chart display caches with their embedded, unchanged worksheets."""
from __future__ import annotations
import io
import math
import posixpath
import zipfile
from decimal import Decimal,InvalidOperation
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from openpyxl.utils.exceptions import InvalidFileException

C='http://schemas.openxmlformats.org/drawingml/2006/chart'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def refresh_numeric_caches(data: bytes) -> tuple[bytes,dict]:
    replacements={};updated=0;charts=0;unresolved=[]
    parser=etree.XMLParser(resolve_entities=False,no_network=True)
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        for name in package.namelist():
            if not name.startswith('ppt/charts/') or not name.endswith('.xml') or '/_rels/' in name:continue
            root=etree.fromstring(package.read(name),parser)
            external=root.find('{'+C+'}externalData')
            if external is None:continue
            relpath=posixpath.join(posixpath.dirname(name),'_rels',posixpath.basename(name)+'.rels')
            if relpath not in package.namelist():continue
            rels=etree.fromstring(package.read(relpath),parser)
            rel=next((r for r in rels if r.get('Id')==external.get('{'+R+'}id')),None)
            if rel is None or rel.get('TargetMode')=='External':continue
            target=rel.get('Target','')
            target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(name),target))
            if target not in package.namelist():continue
            try:workbook=load_workbook(io.BytesIO(package.read(target)),read_only=True,data_only=True,keep_links=False)
            except (ValueError,KeyError,zipfile.BadZipFile,InvalidFileException):
                unresolved.append(name);continue
            changed=False;charts+=1
            try:
                for ref in root.iter('{'+C+'}numRef'):
                    formula=ref.find('{'+C+'}f');cache=ref.find('{'+C+'}numCache')
                    if formula is None or cache is None or not formula.text:continue
                    try:sheet,(x1,y1,x2,y2)=range_to_tuple(formula.text)
                    except ValueError:unresolved.append(name+': range');continue
                    if sheet not in workbook.sheetnames or (x2-x1+1)*(y2-y1+1)>10000:
                        unresolved.append(name+': range');continue
                    values=[cell.value for row in workbook[sheet].iter_rows(min_row=y1,max_row=y2,min_col=x1,max_col=x2) for cell in row]
                    for point in cache.findall('{'+C+'}pt'):
                        index=int(point.get('idx','-1'));value=point.find('{'+C+'}v')
                        if value is None or not 0<=index<len(values):continue
                        actual=values[index]
                        if isinstance(actual,bool) or not isinstance(actual,(int,float)) or not math.isfinite(actual):continue
                        try:same=Decimal(value.text or '')==Decimal(str(actual))
                        except InvalidOperation:same=False
                        if not same:value.text=str(actual);updated+=1;changed=True
                if changed:replacements[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            finally:workbook.close()
        if not replacements:return data,{'charts':charts,'updated_points':0,'unresolved':unresolved}
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as result:
            for member in package.infolist():result.writestr(member,replacements.get(member.filename,package.read(member.filename)))
    return out.getvalue(),{'charts':charts,'updated_points':updated,'unresolved':unresolved}
