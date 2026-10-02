"""Text extraction from source documents: txt, md, pdf, docx, pptx (text only, size limits, no macros or scripts run).

Text decoding, the zip-bomb guard and the Word reader are the shared ones of Hoard Link (``hoard_link.docs``); a PDF
without a text layer (a scan) is handed to the family document service (Kafka's OCR) when it is running.
"""

from __future__ import annotations

import io
import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePath
from typing import Any

from .errors import CiceroError
from .hoard_link import fam_docs
from .hoard_link.docs import readers_lite, textclean

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 400_000
MAX_PDF_PAGES = 400
SUFFIX_KIND = {".txt": "text", ".text": "text", ".md": "markdown", ".markdown": "markdown", ".pdf": "pdf", ".docx": "docx", ".pptx": "pptx"}


@dataclass
class Extracted:
    title: str
    kind: str
    text: str
    truncated: bool = False


def kind_for(filename: str) -> str:
    suffix = PurePath(filename or "").suffix.lower()
    if suffix not in SUFFIX_KIND:
        raise CiceroError(f"Unsupported file type {suffix or '(no extension)'}. Use: {', '.join(sorted(SUFFIX_KIND))}.")
    return SUFFIX_KIND[suffix]


def decode_text(data: bytes) -> str:
    """Bytes to text (BOM, UTF-8, Windows-1252, Latin-1); a file with NUL bytes is not text."""
    try:
        return textclean.decode_text(data, reject_binary=True)
    except ValueError as error:
        raise CiceroError("The file looks binary, not text.") from error


def _check_zip(data: bytes, required: str) -> None:
    """The container checks of an Office file: a zip, with the part that makes it that kind of document, that is not a zip bomb."""
    if data[:2] != b"PK":
        raise CiceroError("The file is not a valid Office document (it is not a zip container).")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if required not in z.namelist():
                raise CiceroError("The file is not a valid Office document of that type.")
            readers_lite.check_zip(z)
    except zipfile.BadZipFile as error:
        raise CiceroError("The file is damaged: it cannot be opened as a zip container.") from error
    except readers_lite.ZipBombError as error:
        raise CiceroError(f"The document was refused: {error}.") from error


def _scan_text(data: bytes, filename: str) -> str:
    """The text of a scanned PDF from the family document service (OCR); empty when nobody can read it (the service is not running)."""
    try:
        with tempfile.TemporaryDirectory(prefix="cicero-ocr-") as folder:
            path = Path(folder) / (PurePath(filename).stem[:60] or "scan").replace(" ", "_")
            path = path.with_suffix(".pdf")
            path.write_bytes(data)
            got = fam_docs.extract(str(path), ocr="auto", local_fallback=False)
    except Exception:  # noqa: BLE001 - OCR is an extra: a failure leaves the file "without readable text"
        return ""
    return str(got.get("text") or "") if got.get("ok") else ""


def _pdf(data: bytes, filename: str = "scan.pdf") -> Extracted:
    if data[:5] != b"%PDF-":
        raise CiceroError("The file is not a PDF.")
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover
        raise CiceroError("Reading PDFs needs pypdf (pip install pypdf).") from error
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise CiceroError("The PDF is password protected.")
            except CiceroError:
                raise
            except Exception as error:  # noqa: BLE001
                raise CiceroError("The PDF is encrypted and cannot be read.") from error
        pages = [(page.extract_text() or "").strip() for page in reader.pages[:MAX_PDF_PAGES]]
        meta = reader.metadata
    except CiceroError:
        raise
    except Exception as error:  # noqa: BLE001 - pypdf raises many types on damaged files
        raise CiceroError(f"The PDF could not be read: {error}") from error
    title = str(getattr(meta, "title", "") or "") if meta else ""
    text = "\n\n".join(t for t in pages if t)
    if not text.strip():  # a scan: the family OCR reads it when Kafka's Hoard is running
        text = _scan_text(data, filename)
    return Extracted(title.strip()[:300], "pdf", text)


def _docx(data: bytes) -> Extracted:
    _check_zip(data, "word/document.xml")
    got = readers_lite.read_any("document.docx", data)
    if got.get("error") or got.get("kind") != "docx":
        raise CiceroError(f"The DOCX could not be read: {got.get('error') or 'it is not a Word document'}")
    lines: list[str] = []
    for unit in got["units"]:  # one unit per heading: "## heading" then its paragraphs, "- " items and " | " table rows
        if unit["title"]:
            lines.append(f"## {unit['title']}")
        if unit["text"]:
            lines.append(unit["text"])
    title = "" if got["title"] == "document" else str(got["title"] or "")  # read_any falls back to the (made-up) name
    return Extracted(title.strip()[:300], "docx", "\n\n".join(lines))


def _shape_texts(shape: Any) -> list[str]:
    out: list[str] = []
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # group
        for sub in shape.shapes:
            out += _shape_texts(sub)
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if text:
                out.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return out


def _pptx(data: bytes) -> Extracted:
    _check_zip(data, "ppt/presentation.xml")
    try:
        from pptx import Presentation
    except ImportError as error:  # pragma: no cover
        raise CiceroError("Reading PPTX needs python-pptx.") from error
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as error:  # noqa: BLE001
        raise CiceroError(f"The PPTX could not be read: {error}") from error
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        title_shape = slide.shapes.title if slide.shapes.title is not None else None
        heading = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            texts += _shape_texts(shape)
        block = [f"## {heading or f'Slide {i}'}"] + texts
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                block.append(f"(notes) {notes}")
        parts.append("\n".join(block))
    title = (prs.core_properties.title or "").strip()
    return Extracted(title[:300], "pptx", "\n\n".join(parts))


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_bytes(data: bytes, filename: str, *, max_bytes: int = MAX_FILE_BYTES) -> Extracted:
    if len(data) > max_bytes:
        raise CiceroError(f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    if not data:
        raise CiceroError("The file is empty.")
    kind = kind_for(filename)
    if kind == "pdf":
        result = _pdf(data, filename)
    elif kind == "docx":
        result = _docx(data)
    elif kind == "pptx":
        result = _pptx(data)
    else:
        result = Extracted("", kind, decode_text(data))
    text = normalize_text(result.text)
    if not text:
        hint = " A scanned PDF needs OCR first." if kind == "pdf" else ""
        raise CiceroError(f"The file has no readable text.{hint}")
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    title = result.title or PurePath(filename).stem
    return Extracted(title[:300], result.kind, text, truncated)
