"""Interactive, editable scene overlays on source-template furniture."""
from __future__ import annotations
import base64
import hashlib
import io
import json
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import MSO_COLOR_TYPE
from pptx.oxml.ns import qn
from .layouts import Box,Para,build_plan,fit_size,Ctx
from .templates import TemplateBinding,PX,is_slide_number
from .template_pptx import _template_theme
from .themes import theme_of
from .render import CSS,esc,_theme_vars,_asset_src,motion_box_html,present_document,SCRIPT,MEDIA_SCRIPT,linked_attributes
from . import native_slides


def scenes(deck,assets):
    descriptor=deck['template'];prs=Presentation(descriptor['path']);result=[]
    for index,item in enumerate(deck['slides'],1):
        b=TemplateBinding.model_validate(descriptor['bindings'].get(item['layout']) or descriptor['bindings']['bullets'])
        source=prs.slides[b.source_slide-1];theme=_template_theme(source,theme_of(deck))
        title=next((s for s in source.shapes if s.shape_id==b.title_shape),None)
        subtitle=next((s for s in source.shapes if s.shape_id==b.subtitle_shape),None)
        slots=[s for s in source.shapes if (s.has_text_frame and s.text.strip() or s.has_chart or s.has_table)
            and s.shape_id not in {b.title_shape,b.subtitle_shape,*b.keep_text_shapes} and not is_slide_number(s)]
        bodies=[s for s in slots if s.has_text_frame and s.text.strip()]
        if bodies:
            body=max(bodies,key=lambda s:len(s.text));runs=[r for p in body.text_frame.paragraphs for r in p.runs if r.text.strip()]
            if runs:
                if runs[0].font.name:theme['fonts']['body']=runs[0].font.name
                if runs[0].font.color.type==MSO_COLOR_TYPE.RGB:theme['colors']['text']='#'+str(runs[0].font.color.rgb)
        if b.text_color:theme['colors']['text']=b.text_color
        if b.muted_color:theme['colors']['muted']=b.muted_color
        if b.chart_font_points is not None:theme['chart_font_points']=b.chart_font_points
        plan=build_plan(item,deck_title=deck['title'],theme=theme,number=index,total=len(deck['slides']),lang=deck['language'],assets=assets)
        boxes=[box for box in plan.boxes if box.kind!='rect' and box.role not in ('footer','title')]
        if subtitle is not None:boxes=[box for box in boxes if [p.text for p in box.paras]!=[item.get('subtitle') or '']]
        if b.content_frame:frame=b.content_frame
        elif slots:
            x=min(s.left for s in slots)/PX;y=min(s.top for s in slots)/PX
            frame=[x,y,max(s.left+s.width for s in slots)/PX-x,max(s.top+s.height for s in slots)/PX-y]
        else:frame=[prs.slide_width/PX*.06,prs.slide_height/PX*.25,prs.slide_width/PX*.88,prs.slide_height/PX*.6]
        moved=[]
        if boxes:
            bx=min(s.x for s in boxes);by=min(s.y for s in boxes);bw=max(s.x+s.w for s in boxes)-bx;bh=max(s.y+s.h for s in boxes)-by
            fx,fy,fw,fh=frame;sx=fw/max(1,bw);sy=fh/max(1,bh);scale=min(sx,sy)
            moved=[replace(box,x=fx+(box.x-bx)*sx,y=fy+(box.y-by)*sy,w=box.w*sx,h=box.h*sy,
                paras=[replace(p,size=p.size*scale,indent=round(p.indent*scale),space_after=p.space_after*scale) for p in box.paras]) for box in boxes]
        for shape,text,override,points,role in [(title,item['title'],b.title_frame,b.title_font_points,'title'),
            (subtitle,item.get('subtitle') or '',b.subtitle_frame,None,'subtitle')]:
            if shape is None or not text:continue
            runs=[r for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()]
            run=runs[-1] if runs else None
            family=run.font.name if run and run.font.name and not run.font.name.startswith('+') else theme['fonts']['heading']
            theme['fonts']['heading']=family
            size=(points/.75 if points else next((r.font.size.pt/.75 for r in reversed(runs) if r.font.size),48))
            x,y,w,h=override or [v/PX for v in (shape.left,shape.top,shape.width,shape.height)]
            value=(b.prefix+'\n' if b.prefix and role=='title' else '')+text
            fitted,_=fit_size(lambda n:[Para(value,n)],w-20,h-10,round(size),18,Ctx(theme));size=min(size,fitted)
            color='#'+str(run.font.color.rgb) if run and run.font.color.type==MSO_COLOR_TYPE.RGB else theme['colors']['accent2']
            if run and run.font.color.type==MSO_COLOR_TYPE.SCHEME:
                scheme=run._r.rPr.find('.//'+qn('a:schemeClr'))
                if scheme is not None:
                    slot=scheme.get('val');slot={'bg1':'lt1','bg2':'lt2','tx1':'dk1','tx2':'dk2'}.get(slot,slot)
                    color=theme.get('source_colors',{}).get(slot,color)
            key='title_ink' if role=='title' else 'subtitle_ink';theme['colors'][key]=color
            alignment='center' if item['layout'] in ('title','closing') else 'left'
            moved.append(Box('text',x+10,y+5,w-20,h-10,paras=[Para(value,size,bold=bool(run.font.bold) if run else role=='title',font='heading',color=key,align=alignment,line=1.2)],anchor='middle' if item['layout'] in ('title','closing') else 'top',role=role))
        result.append({'slide':item,'boxes':moved,'theme':theme,'overflow':plan.overflow_where})
    return result,prs.slide_width/PX,prs.slide_height/PX


def document(svc,deck,*,slide_id=None,mode='stacked',inline=True):
    background=deepcopy(deck)
    background['id']='template-furniture'
    background['updated_at']='';background['created_at']=''
    background['sources']=[];background['outline']=[]
    for i,s in enumerate(background['slides']):
        s.update(id=f'furniture-{i}',title='',subtitle='',blocks=[],notes='',revision=0,status='draft',sources=[],created_at='',updated_at='')
    for binding in background['template']['bindings'].values():binding['prefix']=''
    folder,report=svc._native_snapshot(background)
    scene,width,height=scenes(deck,svc.assets_lookup());sections=[]
    src=_asset_src(svc.assets_lookup(),True)
    for index,item in enumerate(scene):
        s=item['slide']
        if slide_id and s['id']!=slide_id:continue
        raw=base64.b64encode((folder/f'slide-{index+1}.png').read_bytes()).decode('ascii')
        theme=item['theme'];vars=_theme_vars(theme)+''.join(f'--{key}:{value};' for key,value in theme['colors'].items() if key.endswith('_ink'))
        content=''.join(motion_box_html(box,theme,deck['language'],svc.assets_lookup(),src) for box in item['boxes'])
        from .transitions import attributes
        sections.append(f'<section class="slide" id="s-{esc(s["id"])}" data-index="{index}"{attributes(s)} aria-label="{esc(s["title"])}"><div class="stage" style="{esc(vars)};width:{width}px;height:{height}px"><img class="furniture" alt="" src="data:image/png;base64,{raw}"/>{content}</div></section>')
    style=CSS+f'.furniture{{position:absolute;inset:0;width:100%;height:100%}}.stage{{width:{width}px;height:{height}px}}.slide{{aspect-ratio:{width}/{height}}}'
    script=SCRIPT.replace('/1280','/'+str(width))
    result=f'<!doctype html><html lang="{deck["language"]}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(deck["title"])}</title><style>{style}</style></head><body class="mode-{mode}" data-mode="{mode}"{linked_attributes(deck,inline)}><main class="deck">'+''.join(sections)+f'</main><script>{script}{MEDIA_SCRIPT.removeprefix("\n<script>\n").removesuffix("</script>\n")}</script></body></html>'
    return present_document(result,width,height) if mode=='present' else result
