"""Native PPTX export (python-pptx): real text boxes, bullet paragraphs, speaker notes, pictures and native charts.

Geometry, fitted font sizes and colours come from :mod:`layouts`, the same boxes the HTML renderer draws, so the
file looks like the preview. Nothing is rasterised: every piece of text stays editable.
"""

from __future__ import annotations

import io
from typing import Any, Callable, Optional

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .layouts import Box, Para, build_plan, is_card
from .themes import hex_to_rgb, series_colors, theme_of

PX = 9525  # EMU per CSS pixel at 96 dpi
SLIDE_W, SLIDE_H = 1280 * PX, 720 * PX  # 13.333 x 7.5 in
AssetLookup = Callable[[str], Optional[dict[str, Any]]]
LANG_IDS = {"es": "es-ES", "en": "en-US"}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor(*hex_to_rgb(hex_color))


def _emu(px: float) -> Emu:
    return Emu(int(round(px * PX)))


def _color(theme: dict[str, Any], key: str) -> str:
    return theme["colors"].get(key, theme["colors"]["text"])


def _set_bullet(paragraph: Any, indent_px: int, color_hex: str) -> None:
    """A real bullet: hanging indent plus <a:buChar>, in the schema order of a:pPr children."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent_px * PX)))
    pPr.set("indent", str(-int(indent_px * PX)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", color_hex.lstrip("#").upper())
    bu_font = etree.SubElement(pPr, qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    # schema order: lnSpc, spcBef, spcAft, buClr, buFont, buChar ... move spacing before the bullet elements
    for tag in ("a:spcAft", "a:spcBef", "a:lnSpc"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
            pPr.insert(0, el)


def _no_bullet(paragraph: Any) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def _fill_text(tf: Any, box: Box, theme: dict[str, Any], lang: str) -> None:
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(box.anchor, MSO_ANCHOR.TOP)
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, p in enumerate(box.paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = p.text
        font = run.font
        font.size = Pt(p.size * 0.75)
        font.bold = p.bold
        font.italic = p.italic
        font.name = theme["fonts"]["heading" if p.font == "heading" else "body"]
        font.color.rgb = _rgb(_color(theme, p.color))
        try:
            from pptx.enum.lang import MSO_LANGUAGE_ID

            font.language_id = MSO_LANGUAGE_ID.SPANISH if lang == "es" else MSO_LANGUAGE_ID.ENGLISH_US
        except Exception:  # noqa: BLE001 - language tagging is cosmetic
            pass
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(p.align, PP_ALIGN.LEFT)
        para.line_spacing = round(p.line / 1.2, 3)
        para.space_after = Pt(p.space_after * 0.75)
        para.space_before = Pt(0)
        if p.bullet:
            _set_bullet(para, p.indent, _color(theme, "accent"))
        else:
            _no_bullet(para)


def _add_chart(slide: Any, box: Box, theme: dict[str, Any], lang: str) -> None:
    block = box.chart or {}
    kind = block.get("chart", "bar")
    cats = list(block.get("categories") or [])
    data = CategoryChartData()
    data.categories = cats
    series = list(block.get("series") or [])
    if kind == "pie":
        series = series[:1]
    for s in series:
        vals = list(s.get("values") or [])[: len(cats)]
        vals += [None] * (len(cats) - len(vals))
        data.add_series(s.get("name") or "", vals)
    ctype = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    frame = slide.shapes.add_chart(ctype, _emu(box.x), _emu(box.y), _emu(box.w), _emu(box.h), data)
    chart = frame.chart
    values = [v for s in series for v in (s.get("values") or []) if isinstance(v, (int, float))]
    # Grouped thousands; the viewer's locale draws the separator (a point in Spanish), as the preview does.
    number_format = "#,##0" if all(float(v).is_integer() for v in values) else "#,##0.0#"
    text_hex, muted_hex = _color(theme, "text"), _color(theme, "muted")
    palette = series_colors(theme)
    chart.font.size = Pt(11)
    chart.font.name = theme["fonts"]["body"]
    chart.font.color.rgb = _rgb(text_hex)
    chart.has_title = False
    plot = chart.plots[0]
    if kind == "pie":
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = _rgb(text_hex)
        plot.vary_by_categories = True
        ser = plot.series[0]
        for i in range(len(cats)):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        plot.data_labels.font.size = Pt(11)
        plot.data_labels.font.color.rgb = _rgb(text_hex)
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.number_format = number_format
        plot.data_labels.number_format_is_linked = False
    else:
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(12)
            chart.legend.font.color.rgb = _rgb(text_hex)
        for i, ser in enumerate(plot.series):
            col = _rgb(palette[i % len(palette)])
            if kind == "bar":
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = col
            else:
                ser.format.line.color.rgb = col
                ser.format.line.width = Pt(2.5)
                ser.smooth = False
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = col
                ser.marker.format.line.color.rgb = col
        if len(cats) * len(series) <= 14:
            plot.has_data_labels = True
            plot.data_labels.show_value = True
            plot.data_labels.font.size = Pt(11)
            plot.data_labels.font.color.rgb = _rgb(text_hex)
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind == "bar" else XL_LABEL_POSITION.ABOVE
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
        cat_axis, val_axis = chart.category_axis, chart.value_axis
        val_axis.tick_labels.number_format = number_format
        val_axis.tick_labels.number_format_is_linked = False
        for axis in (cat_axis, val_axis):
            axis.tick_labels.font.size = Pt(11)
            axis.tick_labels.font.color.rgb = _rgb(muted_hex if axis is val_axis else text_hex)
            axis.format.line.color.rgb = _rgb(muted_hex)
        val_axis.has_major_gridlines = True
        val_axis.major_gridlines.format.line.color.rgb = _rgb(muted_hex)
        val_axis.major_gridlines.format.line.width = Pt(0.5)
        # gridline transparency: keep the axis calm without adding colour
        ln = val_axis.major_gridlines.format.line._get_or_add_ln()
        for clr in ln.iter(qn("a:srgbClr")):
            alpha = etree.SubElement(clr, qn("a:alpha"))
            alpha.set("val", "30000")


def _add_box(slide: Any, box: Box, theme: dict[str, Any], lang: str, assets: AssetLookup, title_shape: Any) -> Any:
    if box.kind == "rect":
        radius = int(theme.get("radius") or 0) if is_card(box) else 0
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, _emu(box.x), _emu(box.y), _emu(box.w), _emu(box.h))
        if radius:  # the adjustment is the corner radius as a share of the shorter side
            shape.adjustments[0] = min(0.5, radius / max(1.0, min(box.w, box.h)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(_color(theme, box.fill or "surface"))
        shape.line.fill.background()
        shape.shadow.inherit = False
        shape.name = "Decoration"
        return shape
    if box.kind == "text":
        if box.role == "title" and title_shape is not None:
            shape = title_shape
            shape.left, shape.top, shape.width, shape.height = _emu(box.x), _emu(box.y), _emu(box.w), _emu(box.h)
            tf = shape.text_frame
        else:
            shape = slide.shapes.add_textbox(_emu(box.x), _emu(box.y), _emu(box.w), _emu(box.h))
            shape.name = "Footer" if box.role == "footer" else "Text"
            tf = shape.text_frame
        _fill_text(tf, box, theme, lang)
        return shape
    if box.kind == "image" and box.image:
        info = assets(box.image["asset_id"])
        if not info:
            return None
        source: Any = str(info["path"])
        if info.get("mime") == "image/webp":  # python-pptx cannot embed WEBP: convert to PNG
            from PIL import Image

            buf = io.BytesIO()
            Image.open(info["path"]).save(buf, format="PNG")
            buf.seek(0)
            source = buf
        pic = slide.shapes.add_picture(source, _emu(box.x), _emu(box.y), _emu(box.w), _emu(box.h))
        alt = box.image.get("alt") or ""
        pic._element.nvPicPr.cNvPr.set("descr", alt)
        pic.name = "Picture"
        return pic
    if box.kind == "chart" and box.chart:
        _add_chart(slide, box, theme, lang)
        return None
    return None


def build_pptx(deck: dict[str, Any], *, assets: AssetLookup) -> bytes:
    theme = theme_of(deck)
    lang = deck.get("language", "es")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    props = prs.core_properties
    props.title = deck["title"]
    props.author = "Cicero's Hoard"
    props.language = LANG_IDS.get(lang, "es-ES")
    props.subject = (deck.get("brief") or "")[:250]
    slides = deck.get("slides", [])
    title_only = prs.slide_layouts[5]
    blank = prs.slide_layouts[6]
    for index, s in enumerate(slides):
        plan = build_plan(s, deck_title=deck["title"], theme=theme, number=index + 1, total=len(slides), lang=lang, assets=assets)
        has_title = any(b.role == "title" for b in plan.boxes)
        slide = prs.slides.add_slide(title_only if has_title else blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(theme["colors"]["background"])
        title_shape = slide.shapes.title if has_title else None
        decorations = []
        for box in plan.boxes:
            shape = _add_box(slide, box, theme, lang, assets, title_shape if box.role == "title" else None)
            if box.kind == "rect" and shape is not None:
                decorations.append(shape)
        # decorations go to the back, behind the title placeholder that the layout created first
        tree = slide.shapes._spTree
        for shape in reversed(decorations):
            tree.remove(shape._element)
            tree.insert(2, shape._element)
        notes = (s.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
