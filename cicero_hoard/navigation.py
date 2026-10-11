"""Interactive object destinations shared by native exports, the editor and MCP."""
import base64
import json
from copy import deepcopy
from typing import Literal,Optional
from pydantic import BaseModel,Field
from . import decks
from .models import NavigationAction
from .errors import CiceroError,RevisionConflict

MARKER='|Cicero-action:'

class NavigationArgs(BaseModel):
    deck_id:str=Field(...,max_length=40)
    action:Literal['inspect','set','clear']='inspect'
    slide_id:Optional[str]=Field(None,max_length=40)
    block_index:Optional[int]=Field(None,ge=0,le=7)
    target:Optional[NavigationAction]=None
    expected_revision:Optional[int]=Field(None,ge=1)

def inspect(deck):
    ids={s['id'] for s in deck['slides']}
    links=[]
    for slide in deck['slides']:
        for index,block in enumerate(slide['blocks']):
            target=block.get('action')
            if target:
                links.append({'slide_id':slide['id'],'block_index':index,'object_id':block.get('id'),
                              'target':target,'valid':target['kind']!='slide' or target['slide_id'] in ids})
    return {'schema_version':'cicero.navigation.v1','links':links,
            'slides':[{'id':s['id'],'title':s['title'],'position':s['position']} for s in deck['slides']],
            'export_notes':{'web':'Nested return restores origin slide and click stage.',
                            'pptx':'Native slide and URL links; return uses last slide viewed, without stage restoration.',
                            'pdf':'Static content, not an interactive route.'}}

def run(svc,args):
    with svc.db.transaction():
        deck=decks.deck_view(svc,args.deck_id)
        if args.action=='inspect':return inspect(deck)
        if args.slide_id is None or args.block_index is None:raise CiceroError('Choose a slide and block_index')
        slide=decks.get_slide(svc,args.deck_id,args.slide_id)
        if args.expected_revision is not None and args.expected_revision!=slide['revision']:
            raise RevisionConflict('The slide changed. Read its current revision before editing its link.')
        if args.block_index>=len(slide['blocks']):raise CiceroError('Choose an existing block')
        if args.action=='set':
            if args.target is None:raise CiceroError('Supply the navigation target')
            if args.target.kind=='slide' and args.target.slide_id not in {s['id'] for s in deck['slides']}:
                raise CiceroError('The destination slide is not in this deck')
        blocks=deepcopy(slide['blocks'])
        if args.action=='clear':blocks[args.block_index].pop('action',None)
        else:blocks[args.block_index]['action']=args.target.model_dump(exclude_none=True)
        if blocks==slide['blocks']:return {**slide,'changed':False}
        return {**decks.patch_slide(svc,args.deck_id,args.slide_id,{'blocks':blocks},reason='interactive object link'),'changed':True}

def stamp(shape,action):
    from pptx.oxml.ns import qn
    node=shape._element.find('.//'+qn('p:cNvPr'))
    encoded=base64.urlsafe_b64encode(json.dumps(action,separators=(',',':')).encode()).decode()
    node.set('descr',(node.get('descr') or '')+MARKER+encoded)

def apply_native(prs,deck):
    from pptx.oxml.ns import qn
    targets={s['id']:prs.slides[i] for i,s in enumerate(deck['slides'])}
    for slide in prs.slides:
        for shape in slide.shapes:
            node=shape._element.find('.//'+qn('p:cNvPr'))
            descr=node.get('descr','') if node is not None else ''
            if MARKER not in descr:continue
            action=json.loads(base64.urlsafe_b64decode(descr.split(MARKER)[-1]))
            kind=action['kind']
            if kind=='slide':
                if action['slide_id'] not in targets:continue
                shape.click_action.target_slide=targets[action['slide_id']]
            elif kind=='url':shape.click_action.hyperlink.address=action['url']
            else:
                jump={'next':'nextslide','previous':'previousslide','return':'lastslideviewed'}[kind]
                node.get_or_add_hlinkClick().set('action','ppaction://hlinkshowjump?jump='+jump)
            if action.get('label'):node.get_or_add_hlinkClick().set('tooltip',action['label'])
