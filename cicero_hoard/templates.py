"""PPTX template inspection and durable attachment, independent of any task."""
from __future__ import annotations

import hashlib
import io
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field

from .errors import CiceroError, Refused
from .models import LAYOUTS
from .hoard_link.atomic import write_bytes_atomic
from .util import jdump

PX = 9525


def is_slide_number(shape):
    return (
        (shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
        or 'slide number placeholder' in shape.name.lower()
        or any(field.get('type') == 'slidenum' for field in shape._element.iter(qn('a:fld')))
    )


class TemplateBinding(BaseModel):
    model_config = ConfigDict(extra='forbid')
    source_slide: int = Field(..., ge=1)
    title_shape: Optional[int] = Field(None, ge=1)
    subtitle_shape: Optional[int] = Field(None, ge=1)
    keep_text_shapes: list[int] = Field(default_factory=list, max_length=100)
    remove_shapes: list[int] = Field(default_factory=list, max_length=100)
    content_frame: Optional[list[float]] = Field(None, min_length=4, max_length=4,
                                                description='Optional x,y,width,height in CSS pixels.')
    prefix: str = Field('', max_length=100)


def read_template(svc, path: str) -> bytes:
    try:p = Path(path).expanduser().resolve(strict=True)
    except OSError:raise CiceroError(f'Template not found: {path}') from None
    if p.suffix.lower() != '.pptx' or not p.is_file():
        raise CiceroError('Choose a .pptx template file.')
    if svc.config.file_roots and not any(p.is_relative_to(r.resolve()) for r in svc.config.file_roots):
        raise Refused('Template is outside CICERO_FILE_ROOTS.')
    if p.stat().st_size > svc.config.max_upload_bytes:raise CiceroError('Template exceeds the configured upload size.')
    return p.read_bytes()


def _presentation(data: bytes):
    try:prs = Presentation(io.BytesIO(data))
    except Exception as e:raise CiceroError(f'Invalid PPTX template: {type(e).__name__}') from e
    if not prs.slides:
        # A master-only template still supplies usable source slides.
        for layout in list(prs.slide_layouts)[:40]:prs.slides.add_slide(layout)
    if not prs.slides:raise CiceroError('The template has no slide examples or layouts.')
    return prs


def shape_info(s):
    text = s.text if s.has_text_frame else ''
    sizes = [r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs if r.font.size] if s.has_text_frame else []
    placeholder = int(s.placeholder_format.type) if s.is_placeholder else None
    return {'id':s.shape_id,'name':s.name,'kind':'chart' if s.has_chart else 'table' if s.has_table else 'text' if s.has_text_frame else 'drawing',
            'placeholder_type':placeholder,'text':text[:240], 'font_points':max(sizes) if sizes else None,
            'frame':[round(v/PX,2) for v in (s.left,s.top,s.width,s.height)]}


def inspect_bytes(data: bytes):
    prs = _presentation(data)
    slides=[]
    for n,s in enumerate(prs.slides,1):
        slides.append({'number':n,'layout':s.slide_layout.name,'shapes':[shape_info(x) for x in s.shapes]})
    return {'sha256':hashlib.sha256(data).hexdigest(),'width':prs.slide_width/PX,'height':prs.slide_height/PX,
            'width_emu':int(prs.slide_width),'height_emu':int(prs.slide_height),
            'masters':len(prs.slide_masters),'layouts':[l.name for l in prs.slide_layouts],'slides':slides}


def infer_binding(slide, number: int, role: str):
    texts=[s for s in slide.shapes if s.has_text_frame and s.text.strip()
           and not is_slide_number(s)]
    title=slide.shapes.title
    if title is None and texts:
        title=max(texts,key=lambda s: (shape_info(s)['font_points'] or 20, -(s.top)))
    others=[s for s in texts if title is None or s.shape_id!=title.shape_id]
    subtitle=next((s for s in others if s.is_placeholder and s.placeholder_format.type==PP_PLACEHOLDER.SUBTITLE),None)
    if role=='title' and subtitle is None and others:subtitle=min(others,key=lambda s:s.top)
    return TemplateBinding(source_slide=number,title_shape=title.shape_id if title is not None else None,
                           subtitle_shape=subtitle.shape_id if subtitle is not None else None)


def default_bindings(prs):
    body=next((i for i,s in enumerate(prs.slides) if any(x.has_text_frame and len(x.text)>200 for x in s.shapes)),
              min(1,len(prs.slides)-1))
    return {role:infer_binding(prs.slides[0 if role=='title' else len(prs.slides)-1 if role=='closing' else body],
                               1 if role=='title' else len(prs.slides) if role=='closing' else body+1,role).model_dump()
            for role in LAYOUTS}


def attach_template(svc, deck_id: str, data: bytes, name: str, bindings=None):
    from .decks import _deck_row, deck_view
    _deck_row(svc,deck_id);prs=_presentation(data);chosen=default_bindings(prs)
    for role, binding in (bindings or {}).items():
        if role not in LAYOUTS:raise CiceroError(f'Unknown Cicero layout {role!r}.')
        b=binding if isinstance(binding,TemplateBinding) else TemplateBinding.model_validate(binding)
        chosen[role]=b.model_dump()
    for role,raw in chosen.items():
        b=TemplateBinding.model_validate(raw)
        if b.source_slide>len(prs.slides):raise CiceroError(f'{role}: source_slide does not exist.')
        ids={s.shape_id for s in prs.slides[b.source_slide-1].shapes}
        selected=[x for x in (b.title_shape,b.subtitle_shape) if x is not None]+b.keep_text_shapes+b.remove_shapes
        if not set(selected)<=ids:raise CiceroError(f'{role}: a shape ID is absent from the selected source slide.')
        if b.title_shape is not None and b.title_shape==b.subtitle_shape:raise CiceroError('Title and subtitle must use different shapes.')
        if b.content_frame:
            x,y,w,h=b.content_frame
            if not (x>=0 and y>=0 and w>0 and h>0 and x+w<=prs.slide_width/PX and y+h<=prs.slide_height/PX):
                raise CiceroError(f'{role}: content_frame is outside the template slide.')
    digest=hashlib.sha256(data).hexdigest();path=svc.config.templates_dir/(digest+'.pptx')
    path.parent.mkdir(parents=True,exist_ok=True)
    if not path.exists():write_bytes_atomic(path,data)
    descriptor={'name':name,'sha256':digest,'path':str(path),'bindings':chosen,
                'width_emu':int(prs.slide_width),'height_emu':int(prs.slide_height),'format':'pptx'}
    svc.db.execute('UPDATE decks SET template=?,updated_ts=? WHERE id=?',(jdump(descriptor),svc.clock(),deck_id))
    return deck_view(svc,deck_id)
