"""Immutable tabular source snapshots and reproducible chart construction."""
from __future__ import annotations
import csv
import io
import json
import math
import re
import statistics
from collections import defaultdict
from datetime import date,datetime
from pathlib import Path
from zipfile import BadZipFile
from typing import Literal,Optional,Any
from pydantic import BaseModel,Field
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from . import decks as store
from .errors import CiceroError,Refused
from .models import ChartBlock,TableBlock
from .util import sha256_hex

NumberFormat=Literal['auto','es','en']
class ImportArgs(BaseModel):
    deck_id:str
    path:str=''
    text:str=''
    title:str=''
    sheet:Optional[str]=None
    header_row:int=Field(1,ge=1)
    delimiter:Optional[Literal[',',';','\t','|']]=None
    number_format:NumberFormat='auto'

class GetArgs(BaseModel):
    deck_id:str
    source_id:int
    offset:int=Field(0,ge=0)
    limit:int=Field(20,ge=1,le=200)

class Filter(BaseModel):
    column:str
    operator:Literal['eq','ne','gt','ge','lt','le','contains','isnull']='eq'
    value:Any=None

class Aggregate(BaseModel):
    column:str
    operation:Literal['sum','mean','median','min','max','count','last']
    name:str=''

class TransformArgs(BaseModel):
    deck_id:str
    source_id:int
    group_by:list[str]=Field(default_factory=list,max_length=3)
    aggregations:list[Aggregate]=Field(default_factory=list,max_length=12)
    filters:list[Filter]=Field(default_factory=list,max_length=12)
    sort_column:Optional[str]=None
    descending:bool=False
    title:str=''

class VisualArgs(BaseModel):
    deck_id:str
    source_id:int
    kind:Literal['chart','table']='chart'
    chart:Literal['bar','bar_horizontal','line','area','pie','donut','scatter','bubble','dot','dumbbell','band','heatmap','waterfall','funnel']='bar'
    category_column:Optional[str]=None
    series_columns:list[str]=Field(default_factory=list,max_length=10)
    x_column:Optional[str]=None
    lower_column:Optional[str]=None
    upper_column:Optional[str]=None
    size_column:Optional[str]=None
    offset:int=Field(0,ge=0)
    limit:int=Field(30,ge=1,le=2000)
    unit:str=''
    title:str=''

def numeric(value,mode='auto'):
    if value is None or value=='':return None
    if isinstance(value,bool):return None
    if isinstance(value,(int,float)):
        return float(value) if math.isfinite(value) else None
    t=str(value).strip().replace('\u2212','-').replace('\u00a0','').replace('\u202f','')
    if not re.fullmatch(r'[+-]?(?:\d[\d.,]*|[.,]\d+)(?:[eE][+-]?\d+)?',t):return None
    if mode=='es':
        if not re.fullmatch(r'[+-]?(?:(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?|,\d+)(?:[eE][+-]?\d+)?',t):return None
        t=t.replace('.','').replace(',','.')
    elif mode=='en':
        if not re.fullmatch(r'[+-]?(?:(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?|\.\d+)(?:[eE][+-]?\d+)?',t):return None
        t=t.replace(',','')
    elif ',' in t and '.' in t:
        return numeric(t,'es' if t.rfind(',')>t.rfind('.') else 'en')
    elif re.fullmatch(r'[+-]?\d{1,3}[.,]\d{3}',t):return None
    else:t=t.replace(',','.')
    try:value=float(t)
    except ValueError:return None
    return value if math.isfinite(value) else None

def _clean(value):
    if isinstance(value,(date,datetime)):return value.isoformat()
    if isinstance(value,float) and not math.isfinite(value):return None
    return value if isinstance(value,(str,int,float,bool)) or value is None else str(value)

def _snapshot(svc,deck_id,title,payload):
    text=json.dumps(payload,ensure_ascii=False,separators=(',',':'),allow_nan=False)
    with svc.db.transaction():
        parent=payload.get('origin',{}).get('derived_from')
        if parent is not None and not svc.db.one("SELECT id FROM sources WHERE id=? AND deck_id=? AND kind='dataset'",(parent,deck_id)):
            raise CiceroError('The parent dataset was removed before this transformation completed.')
        # Check and insert under the same transaction so retries cannot create duplicates.
        old=svc.db.one("SELECT * FROM sources WHERE deck_id=? AND kind='dataset' AND sha256=?",(deck_id,sha256_hex(text)))
        source=store._source_dict(old) if old else store._insert_source(svc,deck_id,title or 'Tabular data','dataset',text)
    return inspect(svc,GetArgs(deck_id=deck_id,source_id=source['id']))|source|{'reused':bool(old)}

def load(svc,deck_id,source_id):
    source=store.get_source(svc,deck_id,source_id)
    if source['kind']!='dataset':raise CiceroError('Choose a tabular dataset source.')
    return json.loads(source['text'])

def import_data(svc,args,data=None,filename=''):
    store.deck_view(svc,args.deck_id)
    origin={};warnings=[]
    if data is None and args.path:
        path=Path(args.path).expanduser().resolve()
        if svc.config.file_roots and not any(r.resolve()==path or r.resolve() in path.parents for r in svc.config.file_roots):raise Refused('The data file is outside the configured source folders.')
        if not path.is_file():raise CiceroError('The data file does not exist.')
        if path.stat().st_size>svc.config.max_upload_bytes:raise CiceroError('The data file exceeds the existing upload limit.')
        data=path.read_bytes();filename=path.name;origin={'path':str(path)}
    if data is None:data=args.text.encode('utf-8');filename='pasted.csv'
    if not data:raise CiceroError('Supply CSV text or a local CSV, TSV or XLSX file.')
    if len(data)>svc.config.max_upload_bytes:raise CiceroError('The data exceeds the existing upload limit.')
    origin.update(filename=filename,sha256=sha256_hex(data))
    if filename.lower().endswith('.xlsx'):
        try:workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        except (BadZipFile,InvalidFileException,ValueError,KeyError) as error:raise CiceroError('The XLSX workbook cannot be read. Export a valid workbook or CSV.') from error
        try:formula_book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        except Exception:
            workbook.close()
            raise
        try:
            sheet=args.sheet or workbook.sheetnames[0]
            if sheet not in workbook.sheetnames:raise CiceroError('Select an existing worksheet: '+', '.join(workbook.sheetnames))
            rows=list(workbook[sheet].iter_rows(min_row=args.header_row,values_only=True))
            for cached,r in zip(rows,formula_book[sheet].iter_rows(min_row=args.header_row)):
                for value,cell in zip(cached,r):
                    if cell.data_type=='f' and value is None:warnings.append('Uncached formula: '+cell.coordinate)
            origin.update(sheet=sheet,sheets=workbook.sheetnames)
        finally:workbook.close();formula_book.close()
    elif Path(filename).suffix.lower() in ('.csv','.tsv','.txt'):
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:text=data.decode('cp1252')
            except UnicodeDecodeError as error:raise CiceroError('The data text is neither UTF-8 nor readable cp1252. Export UTF-8 CSV.') from error
            warnings.append('Decoded legacy text as cp1252.')
        delimiter=args.delimiter
        if not delimiter:
            try:delimiter=csv.Sniffer().sniff(text[:8192],delimiters=',;\t|').delimiter
            except csv.Error:delimiter='\t' if filename.lower().endswith('.tsv') else ','
        rows=list(csv.reader(io.StringIO(text),delimiter=delimiter))[args.header_row-1:];origin['delimiter']=delimiter
    else:raise CiceroError('Use CSV, TSV or XLSX data.')
    rows=[list(r) for r in rows if any(v not in (None,'') for v in r)]
    if len(rows)<2:raise CiceroError('The data needs a header and at least one row.')
    width=max(len(r) for r in rows);header=rows[0]+['']*(width-len(rows[0]));columns=[]
    for i,value in enumerate(header):
        name=(str(value).strip() if value is not None else '') or f'Column {i+1}';original=name;counter=2
        while name in columns:name=f'{original} ({counter})';counter+=1
        columns.append(name)
    body=[[_clean(v) for v in (row+[None]*width)[:width]] for row in rows[1:]]
    for i,name in enumerate(columns):
        if args.number_format=='auto' and any(re.fullmatch(r'[+-]?\d{1,3}[.,]\d{3}',str(r[i] or '')) for r in body):warnings.append('Ambiguous separators in '+name+': select es or en before numeric mapping.')
    return _snapshot(svc,args.deck_id,args.title or filename,{'schema_version':'cicero.dataset.v1','columns':columns,'rows':body,'number_format':args.number_format,'origin':origin,'warnings':list(dict.fromkeys(warnings))})

def inspect(svc,args):
    data=load(svc,args.deck_id,args.source_id);mode=data['number_format'];rows=data['rows']
    profile=[]
    for i,name in enumerate(data['columns']):
        values=[r[i] for r in rows];numbers=[n for v in values if (n:=numeric(v,mode)) is not None]
        profile.append({'name':name,'numeric':len(numbers),'missing':sum(v in (None,'') for v in values),
            'min':min(numbers) if numbers else None,'max':max(numbers) if numbers else None})
    return {'id':args.source_id,'source_id':args.source_id,'columns':data['columns'],'row_count':len(rows),'profile':profile,
        'rows':rows[args.offset:args.offset+args.limit],'offset':args.offset,'has_more':args.offset+args.limit<len(rows),
        'number_format':mode,'origin':data['origin'],'warnings':data.get('warnings',[])}

def _index(data,name):
    try:return data['columns'].index(name)
    except ValueError:raise CiceroError('Unknown column: '+str(name)) from None

def _numbers(data,rows,name):
    index=_index(data,name);result=[]
    for r in rows:
        value=numeric(r[index],data['number_format'])
        if value is None and r[index] not in (None,''):raise CiceroError(f'Column {name} contains a non-numeric or ambiguous value: {str(r[index])[:80]}')
        result.append(value)
    return result

def visual(svc,args):
    data=load(svc,args.deck_id,args.source_id);rows=data['rows'][args.offset:args.offset+args.limit]
    if not rows:raise CiceroError('The selected row range is empty.')
    if args.kind=='table':
        names=args.series_columns or data['columns'];indices=[_index(data,n) for n in names]
        block=TableBlock(headers=names,rows=[['' if r[i] is None else str(r[i]) for i in indices] for r in rows[:30]],title=args.title or None)
    else:
        if not args.category_column or not args.series_columns:raise CiceroError('Choose the category and value columns.')
        cats=['' if r[_index(data,args.category_column)] is None else str(r[_index(data,args.category_column)]) for r in rows]
        series=[]
        for name in args.series_columns:
            item={'name':name,'values':_numbers(data,rows,name)}
            for key,column in [('x_values',args.x_column),('lower',args.lower_column),('upper',args.upper_column),('sizes',args.size_column)]:
                if column:item[key]=_numbers(data,rows,column)
            series.append(item)
        block=ChartBlock(chart=args.chart,categories=cats,series=series,unit=args.unit or None,title=args.title or None)
    return {'block':block.model_dump(exclude_none=True),'source_ids':[args.source_id],'selected_rows':len(rows),
        'warnings':(['Table uses the first 30 selected rows.'] if args.kind=='table' and len(rows)>30 else []),
        'mapping':args.model_dump(exclude={'deck_id'},exclude_none=True)}

def transform(svc,args):
    data=load(svc,args.deck_id,args.source_id);rows=data['rows'];mode=data['number_format']
    if args.group_by and not args.aggregations:raise CiceroError('Grouping requires at least one aggregation.')
    for f in args.filters:
        index=_index(data,f.column)
        def matches(row):
            value=row[index]
            if f.operator=='isnull':return value in (None,'')
            if f.operator=='contains':return str(f.value) in ('' if value is None else str(value))
            left=numeric(value,mode);right=numeric(f.value,mode)
            if f.operator in ('gt','ge','lt','le'):
                if value in (None,''):return False
                if left is None or right is None:raise CiceroError('Numeric comparisons require unambiguous numeric values.')
            if left is None or right is None:left,right=str(value),str(f.value)
            return {'eq':lambda:left==right,'ne':lambda:left!=right,'gt':lambda:left>right,'ge':lambda:left>=right,'lt':lambda:left<right,'le':lambda:left<=right}[f.operator]()
        rows=[r for r in rows if matches(r)]
    if args.aggregations:
        keys=[_index(data,n) for n in args.group_by];groups=defaultdict(list)
        for row in rows:groups[tuple(row[i] for i in keys)].append(row)
        if not keys and not rows:groups[()]=[]
        columns=[*args.group_by,*[a.name or a.operation+' '+a.column for a in args.aggregations]]
        if len(set(columns))!=len(columns):raise CiceroError('Give each aggregate a unique output name.')
        output=[]
        for key,group in groups.items():
            values=list(key)
            for a in args.aggregations:
                if a.operation in ('count','last'):
                    index=_index(data,a.column)
                    observed=[r[index] for r in group if r[index] not in (None,'')]
                    values.append(len(observed) if a.operation=='count' else observed[-1] if observed else None)
                else:
                    numbers=[v for v in _numbers(data,group,a.column) if v is not None]
                    operations={'sum':sum,'mean':statistics.mean,'median':statistics.median,'min':min,'max':max}
                    values.append(operations[a.operation](numbers) if numbers or a.operation=='sum' else None)
            output.append(values)
        rows=output
    else:columns=data['columns']
    if args.sort_column:
        try:index=columns.index(args.sort_column)
        except ValueError:raise CiceroError('Unknown sort column.') from None
        def sort_key(row):
            n=numeric(row[index],mode)
            return (0,n) if n is not None else (1,str(row[index]))
        present=[r for r in rows if r[index] not in (None,'')]
        missing=[r for r in rows if r[index] in (None,'')]
        rows=sorted(present,key=sort_key,reverse=args.descending)+missing
    payload={'schema_version':'cicero.dataset.v1','columns':columns,'rows':rows,'number_format':mode,'warnings':[],
        'origin':{'derived_from':args.source_id,'transformation':args.model_dump(exclude={'deck_id'},exclude_none=True)}}
    return _snapshot(svc,args.deck_id,args.title or 'Derived data',payload)

def track(svc,args,result):
    from .agent_undo import track_source_add
    if result.get('reused'):return {'objects':[],'ids':[],'etag':'reused'}
    return track_source_add(svc,args,result)

def undo(svc,record,dry_run=False):
    from .agent_undo import undo_source_add,_id_of,_conflict
    if record.get('etag')=='reused':return {'unchanged':True}
    deck_id=_id_of(record,'deck_id');source_id=int(_id_of(record,'source_id') or 0)
    with svc.db.transaction():
        children=[row['id'] for row in svc.db.query("SELECT id FROM sources WHERE deck_id=? AND kind='dataset' AND json_extract(text,'$.origin.derived_from')=?",(deck_id,source_id))]
        if children:
            if dry_run and _session_removes_children(svc,record,deck_id,children):
                return {'would_remove':source_id,'after':'derived snapshots from this session are undone first'}
            raise _conflict('Derived datasets still depend on this source. Undo those snapshots first: '+', '.join(map(str,children)))
        return undo_source_add(svc,record,dry_run)

def _session_removes_children(svc,record,deck_id,children):
    from .agent_undo import _id_of
    from .hoard_link.agent_journal import Journal
    rows=Journal(svc.config.data_dir).query(session=str(record.get('session') or ''),kind='write',limit=1000)
    planned={int(_id_of(r,'source_id') or 0) for r in rows if r.get('ok') and not r.get('undone')
        and r.get('tool') in ('dataset_import','dataset_transform') and r.get('etag')!='reused'
        and r.get('id')!=record.get('id') and float(r.get('ts') or 0)>=float(record.get('ts') or 0)
        and _id_of(r,'deck_id')==deck_id}
    return set(children).issubset(planned)
