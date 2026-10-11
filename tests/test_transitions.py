import io
import pytest
from pptx import Presentation
from cicero_hoard import transitions, decks
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.render import render_document, present_document
from cicero_hoard.layouts import build_plan
from cicero_hoard.themes import theme_of
from cicero_hoard.errors import CiceroError
from conftest import new_deck


def test_morph_objects_web_native_and_reset(services):
    deck=new_deck(services)
    for frame in ([0,0,.3,.4],[.4,.2,.5,.6]):
        decks.add_slide(services,deck['id'],title='Matched',layout='freeform',blocks=[{'type':'text','id':'subject','text':'Subject','frame':frame}])
    transitions.run(services,transitions.Args(deck_id=deck['id'],action='set',default={'effect':'morph','duration_ms':900}))
    deck=decks.deck_view(services,deck['id'])
    doc,_=render_document(deck,assets=services.assets_lookup())
    doc=present_document(doc)
    assert doc.count('data-auto-animate data-auto-animate-id="default"')==2
    assert doc.count('data-id="subject-0"')==2
    assert 'autoAnimate:!reduced' in doc
    report=transitions.run(services,transitions.Args(deck_id=deck['id'],action='inspect'))
    assert report['pairs'][0]['matched']==['subject']
    prs=Presentation(io.BytesIO(build_pptx(deck,assets=services.assets_lookup())))
    assert all(any(s.name.startswith('!!subject-0') for s in slide.shapes) for slide in prs.slides)
    assert all('morph' in slide._element.xml for slide in prs.slides)
    transitions.run(services,transitions.Args(deck_id=deck['id'],action='reset'))
    assert not decks.deck_view(services,deck['id'])['transitions']


def test_duplicate_block_ids_and_invalid_override_rejected(services):
    deck=new_deck(services)
    with pytest.raises(CiceroError):decks.add_slide(services,deck['id'],blocks=[{'type':'text','text':'A','id':'same'},{'type':'text','text':'B','id':'same'}])
    with pytest.raises(CiceroError):transitions.run(services,transitions.Args(deck_id=deck['id'],action='set',slides={'missing':{'effect':'morph'}}))


def test_none_disables_internal_stage_effects(services):
    deck=new_deck(services)
    slide=decks.add_slide(services,deck['id'],blocks=[{'type':'timeline','build':1,'effect':'none','items':[{'label':'A'},{'label':'B'}]}])
    plan=build_plan(slide,deck_title=deck['title'],theme=theme_of(deck),number=1,total=1,assets=services.assets_lookup())
    assert all(b.effect=='none' for b in plan.boxes if b.build)
