import io
import shutil
import subprocess
import pytest
from pptx import Presentation
from cicero_hoard import media,decks
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.render import render_document,present_document
from cicero_hoard.models import MediaBlock
from cicero_hoard.errors import CiceroError
from cicero_hoard.presentation import start,StartArgs,media_control,MediaControlArgs
from conftest import new_deck


@pytest.fixture
def clip(tmp_path):
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):pytest.skip('Optional media runtime')
    path=tmp_path/'test.mp4'
    subprocess.run([shutil.which('ffmpeg'),'-v','error','-f','lavfi','-i','testsrc2=size=320x180:rate=10','-f','lavfi','-i','sine=frequency=440','-t','2','-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac',str(path)],check=True,timeout=30)
    return path


def test_import_probe_poster_offline_and_native_clip(services,clip):
    deck=new_deck(services)
    args=media.ImportArgs(deck_id=deck['id'],path=str(clip))
    imported=media.import_media(services,args)
    assert imported['duration']==pytest.approx(2,abs=.1) and imported['kind']=='video'
    assert imported['poster_id'] and imported['width']==320
    assert media.import_media(services,args)['reused']
    block=MediaBlock(type='video',asset_id=imported['asset_id'],start=.5,end=1.5,muted=True,autoplay=True,playback_rate=2,
        subtitles='WEBVTT\n\n00:00.500 --> 00:01.500\nCaption\n').model_dump(exclude_none=True)
    decks.add_slide(services,deck['id'],title='Media',blocks=[block])
    deck=decks.deck_view(services,deck['id'])
    doc,_=render_document(deck,assets=services.assets_lookup(),inline=True)
    doc=present_document(doc)
    assert 'data:video/mp4;base64,' in doc and 'data:text/vtt;base64,' in doc and 'data-start="0.5"' in doc
    assert 'media.playbackRate' in doc and 'data-autoplay' in doc
    prs=Presentation(io.BytesIO(build_pptx(deck,assets=services.assets_lookup())))
    xml=prs.slides[0]._element.xml
    assert 'videoFile' in xml and 'vol="0"' in xml
    assert 'delay="0"' in xml
    # Verify embedded compatibility media, not the untouched original.
    import zipfile
    with zipfile.ZipFile(io.BytesIO(build_pptx(deck,assets=services.assets_lookup()))) as z:
        name=next(n for n in z.namelist() if n.endswith('.mp4'))
        converted=clip.parent/'converted.mp4';converted.write_bytes(z.read(name))
    assert media._probe(converted)['duration']==pytest.approx(.5,abs=.15)
    assert clip.stat().st_size>converted.stat().st_size


def test_native_export_keeps_media_timing_with_morph_and_click_builds(services,clip):
    deck=new_deck(services)
    imported=media.import_media(services,media.ImportArgs(deck_id=deck['id'],path=str(clip)))
    slide=decks.add_slide(services,deck['id'],blocks=[{'type':'video','asset_id':imported['asset_id'],'autoplay':True,'id':'movie'},
        {'type':'text','text':'Reveal','build':1,'effect':'fade','id':'takeaway'}])
    from cicero_hoard.agent_tools import call_tool
    call_tool(services,'deck_transitions',{'deck_id':deck['id'],'action':'set','default':{'effect':'morph'}})
    exported=services.export(deck['id'],'pptx')
    path,_=decks.export_file(services,exported['id'])
    prs=Presentation(str(path));xml=prs.slides[0]._element.xml
    assert '<p:video>' in xml and 'morph' in xml and '!!takeaway' in xml
    from lxml import etree
    from cicero_hoard.animations import NS
    ids=prs.slides[0]._element.xpath('.//p:cTn/@id')
    assert len(ids)==len(set(ids))
    assert exported['animations']['animated']>=1
    decks.add_slide(services,deck['id'],blocks=[{'type':'text','text':'Next'}])
    from cicero_hoard.presentation import control,ControlArgs
    start(services,StartArgs(deck_id=deck['id']))
    command=media_control(services,MediaControlArgs(deck_id=deck['id'],action='play'))
    assert command['media_command']['slide_id']==slide['id']
    navigated=control(services,ControlArgs(deck_id=deck['id'],action='goto',slide_index=1))
    assert navigated['media_command'] is None


def test_audio_embed_and_live_media_commands(services,tmp_path):
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):pytest.skip('Optional media runtime')
    path=tmp_path/'tone.wav'
    subprocess.run([shutil.which('ffmpeg'),'-v','error','-f','lavfi','-i','sine=frequency=220','-t','1',str(path)],check=True,timeout=30)
    deck=new_deck(services)
    imported=media.import_media(services,media.ImportArgs(deck_id=deck['id'],path=str(path)))
    decks.add_slide(services,deck['id'],blocks=[MediaBlock(type='audio',asset_id=imported['asset_id']).model_dump()])
    start(services,StartArgs(deck_id=deck['id']))
    result=media_control(services,MediaControlArgs(deck_id=deck['id'],action='seek',value=.5))
    assert result['media_command']['value']==.5 and 'notes' not in result
    with pytest.raises(CiceroError):media_control(services,MediaControlArgs(deck_id=deck['id'],action='volume',value=2))
    with pytest.raises(CiceroError):media_control(services,MediaControlArgs(deck_id=deck['id'],action='seek',value=10))
    prs=Presentation(io.BytesIO(build_pptx(decks.deck_view(services,deck['id']),assets=services.assets_lookup())))
    assert 'audioFile' in prs.slides[0]._element.xml and '<p:audio>' in prs.slides[0]._element.xml


def test_invalid_ranges_types_and_cross_deck_media(services,clip):
    deck=new_deck(services);other=new_deck(services)
    imported=media.import_media(services,media.ImportArgs(deck_id=deck['id'],path=str(clip)))
    with pytest.raises(CiceroError):decks.add_slide(services,other['id'],blocks=[{'type':'video','asset_id':imported['asset_id']}])
    with pytest.raises(CiceroError):decks.add_slide(services,deck['id'],blocks=[{'type':'audio','asset_id':imported['asset_id']}])
    with pytest.raises(CiceroError):decks.add_slide(services,deck['id'],blocks=[{'type':'video','asset_id':imported['asset_id'],'start':9}])
