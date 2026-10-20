import io
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import pytest
from cicero_hoard import decks, agent_tools
from cicero_hoard.models import DeckCreate
from cicero_hoard.errors import CiceroError
from conftest import png_bytes


def template(path, wide=True):
    p=Presentation();p.slide_width=Inches(16 if wide else 10);p.slide_height=Inches(9 if wide else 7.5)
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('125566')
    banner = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(16 if wide else 10), Inches(.15))
    banner.name = 'Empty decorative banner'
    banner.fill.solid();banner.fill.fore_color.rgb=RGBColor.from_string('FFAA00')
    title=s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2));title.text='Example title';title.text_frame.paragraphs[0].runs[0].font.size=Pt(40)
    title.text_frame.paragraphs[0].runs[0].font.name='Georgia'
    logo=s.shapes.add_picture(io.BytesIO(png_bytes()),Inches(.2),Inches(.2),Inches(.5),Inches(.5))
    footer=s.shapes.add_textbox(Inches(1),Inches(6),Inches(8),Inches(.5));footer.text='Brand footer'
    footer.text_frame.paragraphs[0].runs[0].hyperlink.address='https://example.test/brand'
    body=s.shapes.add_textbox(Inches(1),Inches(3),Inches(8),Inches(2));body.text='Example body '*30
    body.text_frame.paragraphs[0].runs[0].font.name='Tahoma'
    p.save(path)
    return {'source_slide':1,'title_shape':title.shape_id,'keep_text_shapes':[footer.shape_id],
            'content_frame':[96,288,768,192]}

def test_timeline_rules_and_freeform_positions_survive_template_mapping(services,tmp_path):
    from cicero_hoard.template_live import scenes
    file=tmp_path/'composition-template.pptx';binding=template(file)
    deck=decks.create_deck(services,DeckCreate(title='Reusable scene',slide_count=3))
    decks.add_slide(services,deck['id'],layout='freeform',title='Precise placement',blocks=[
        {'type':'text','text':'Right half','frame':[.5,0,.5,.5]}])
    decks.add_slide(services,deck['id'],layout='evidence',title='Milestones',blocks=[
        {'type':'timeline','items':[{'label':'Source'},{'label':'Review'},{'label':'Deliver'}]}])
    agent_tools.call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file),
        'bindings':{'freeform':binding,'evidence':binding}})
    snapshot=decks.deck_view(services,deck['id']);scene,_,_=scenes(snapshot,services.assets_lookup())
    text=next(b for b in scene[0]['boxes'] if b.paras and b.paras[0].text=='Right half')
    assert text.x==pytest.approx(96+768*.5)
    assert sum(b.kind=='rule' for b in scene[1]['boxes'])==2
    from cicero_hoard.export_pptx import build_pptx
    prs=Presentation(io.BytesIO(build_pptx(snapshot,assets=services.assets_lookup())))
    assert sum(s.name=='Decoration' for s in prs.slides[1].shapes)>=2


@pytest.mark.parametrize('wide', [True, False])
def test_template_frame_preserves_generated_picture_aspect(services, tmp_path, wide):
    file=tmp_path/'picture-template.pptx';binding=template(file,wide=wide)
    deck=decks.create_deck(services,DeckCreate(title='Figure proportions',slide_count=3))
    asset=decks.add_asset(services,deck['id'],png_bytes(),'figure.png')
    decks.add_slide(services,deck['id'],title='Figure',layout='image_text',
                    blocks=[{'type':'image','asset_id':asset['asset_id'],'caption':'Caption'},
                            {'type':'bullets','items':['Kept']}])
    agent_tools.call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file),
                         'bindings':{'image_text':binding}})
    result=agent_tools.call_tool(services,'deck_export',{'deck_id':deck['id'],'format':'pptx'})
    picture=next(s for s in Presentation(result['path']).slides[0].shapes if s.name=='Picture')
    assert picture.width/picture.height == pytest.approx(asset['width']/asset['height'],abs=1e-5)


def test_native_roundtrip_preserves_template_and_independent_charts_notes(services,tmp_path):
    file=tmp_path/'brand.pptx';binding=template(file)
    d=decks.create_deck(services,DeckCreate(title='Any topic',slide_count=3))
    for title in ['First new slide','Second new slide']:
        decks.add_slide(services,d['id'],layout='chart',title=title,blocks=[{'type':'chart','chart':'bar','categories':['A','B'],
            'series':[{'name':'Measured','values':[3.49,4.26]}]}],notes='Verified speaker notes')
    inspected=agent_tools.call_tool(services,'template_inspect',{'path':str(file)})
    assert inspected['width_emu']==int(Inches(16))
    agent_tools.call_tool(services,'deck_template',{'deck_id':d['id'],'path':str(file),'bindings':{'chart':binding}})
    export=agent_tools.call_tool(services,'deck_export',{'deck_id':d['id'],'format':'pptx'});data=Path(export['path']).read_bytes()
    p=Presentation(io.BytesIO(data));assert len(p.slides)==2 and p.slide_width==Inches(16)
    for s in p.slides:
        assert s.background.fill.fore_color.rgb==RGBColor.from_string('125566')
        assert any(x.name == 'Empty decorative banner' for x in s.shapes)
        assert 'Brand footer' in [x.text for x in s.shapes if x.has_text_frame]
        assert any(x.shape_type==13 for x in s.shapes)
        assert s.notes_slide.notes_text_frame.text=='Verified speaker notes'
        assert any(x.has_chart for x in s.shapes)
    assert 'First new slide' in [x.text for x in p.slides[0].shapes if x.has_text_frame]
    footer = next(x for x in p.slides[0].shapes if x.has_text_frame and x.text == 'Brand footer')
    assert footer.text_frame.paragraphs[0].runs[0].hyperlink.address == 'https://example.test/brand'
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert len([n for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')])==2
        assert len([n for n in z.namelist() if n.startswith('ppt/embeddings/')])==2
        from lxml import etree
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        for name in z.namelist():
            if name.startswith('ppt/charts/chart') and name.endswith('.xml'):
                xml = etree.fromstring(z.read(name))
                ids = [int(x.get('val')) for x in xml.findall('.//c:axId', ns)]
                refs = [int(x.get('val')) for x in xml.findall('.//c:crossAx', ns)]
                assert ids and all(0 <= value <= 0xFFFFFFFF for value in ids)
                assert set(refs) == set(ids)


def test_snapshot_is_reusable_after_source_changes_and_clear_restores_default(services,tmp_path,monkeypatch):
    file=tmp_path/'source.pptx';template(file,wide=False)
    d=decks.create_deck(services,DeckCreate(title='Reuse',slide_count=3));decks.add_slide(services,d['id'],title='New body',notes='Notes')
    attached=agent_tools.call_tool(services,'deck_template',{'deck_id':d['id'],'path':str(file)})
    before=Path(attached['template']['path']).read_bytes();file.write_bytes(b'changed original')
    exported=agent_tools.call_tool(services,'deck_export',{'deck_id':d['id'],'format':'pptx'});assert Presentation(exported['path']).slide_width==Inches(10)
    assert Path(attached['template']['path']).read_bytes()==before
    from cicero_hoard import native_slides
    monkeypatch.setattr(native_slides, 'available', lambda: False)
    with pytest.raises(CiceroError):services.export(d['id'],'pdf')
    agent_tools.call_tool(services,'deck_template',{'deck_id':d['id'],'clear':True})
    assert 'template' not in decks.deck_view(services,d['id'])
    assert Presentation(agent_tools.call_tool(services,'deck_export',{'deck_id':d['id'],'format':'pptx'})['path']).slide_width==Inches(13.333333333333334)


def test_bad_binding_is_atomic_and_source_style_is_inspected(services,tmp_path):
    file=tmp_path/'source.pptx';template(file)
    d=decks.create_deck(services,DeckCreate(title='Invalid binding',slide_count=3));before=decks.deck_view(services,d['id'])
    with pytest.raises(CiceroError):agent_tools.call_tool(services,'deck_template',{'deck_id':d['id'],'path':str(file),'bindings':{'bullets':{'source_slide':999}}})
    assert decks.deck_view(services,d['id'])==before


def test_multipart_template_route(client,services,deck,tmp_path):
    file=tmp_path/'brand.pptx';template(file)
    response=client.post(f"/api/decks/{deck['id']}/template",files={'file':('brand.pptx',file.read_bytes(),'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
    assert response.status_code==200 and response.json()['template']['name']=='brand.pptx'
    response=client.get(f"/api/decks/{deck['id']}/template")
    assert response.status_code==200 and response.json()['template']['sha256']


def test_blank_template_still_exports_the_authored_title(services, tmp_path):
    file = tmp_path / 'blank.pptx'
    p = Presentation(); p.slides.add_slide(p.slide_layouts[6]); p.save(file)
    d = decks.create_deck(services, DeckCreate(title='Another topic', slide_count=3))
    decks.add_slide(services, d['id'], layout='title', title='Authored title', subtitle='Authored subtitle')
    agent_tools.call_tool(services, 'deck_template', {'deck_id':d['id'], 'path':str(file)})
    out = agent_tools.call_tool(services, 'deck_export', {'deck_id':d['id'], 'format':'pptx'})
    exported = Presentation(out['path'])
    texts = [s.text for s in exported.slides[0].shapes if s.has_text_frame]
    assert 'Authored title' in texts and 'Authored subtitle' in texts


def test_snapshot_survives_a_moved_data_folder(services,tmp_path):
    # The stored path is absolute (for example a Windows path from another machine); the
    # content-addressed snapshot is still found by name in this installation's templates folder.
    import json
    file=tmp_path/'source.pptx';template(file)
    d=decks.create_deck(services,DeckCreate(title='Moved',slide_count=3));decks.add_slide(services,d['id'],title='Body',notes='Notes')
    attached=agent_tools.call_tool(services,'deck_template',{'deck_id':d['id'],'path':str(file)})['template']
    stored=dict(attached,path='C:\\Users\\someone\\Cicero\\data\\templates\\'+attached['sha256']+'.pptx')
    services.db.execute('UPDATE decks SET template=? WHERE id=?',(json.dumps(stored),d['id']))
    view=decks.deck_view(services,d['id'])
    snapshot=services.config.templates_dir/(attached['sha256']+'.pptx')
    assert Path(view['template']['path'])==snapshot
    exported=agent_tools.call_tool(services,'deck_export',{'deck_id':d['id'],'format':'pptx'})
    assert Presentation(exported['path']).slide_width==Inches(16)
    # a snapshot that is really gone still fails clearly
    snapshot.unlink()
    with pytest.raises(CiceroError):agent_tools.call_tool(services,'deck_export',{'deck_id':d['id'],'format':'pptx'})
