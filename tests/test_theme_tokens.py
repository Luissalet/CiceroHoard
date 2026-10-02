"""Themes made from a design system: the mapping, the contrast fit, fonts, radius, the tool, the routes and the renderers."""
from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from cicero_hoard import agent_tools as at
from cicero_hoard import custom_themes, decks as store, themes
from cicero_hoard.errors import CiceroError
from cicero_hoard.models import DeckPatch
from conftest import FakeHub, make_services, new_deck, sample_roles


def test_font_mapping_keeps_windows_fonts_and_replaces_the_rest():
    assert themes.pick_font("Georgia, 'Times New Roman', serif") == ("Georgia", None)
    name, warning = themes.pick_font("Inter, 'Segoe UI', system-ui, sans-serif")
    assert name == "Segoe UI" and "Inter" in warning
    assert themes.pick_font("'Playfair Display', serif")[0] == "Georgia"
    assert themes.pick_font("ui-monospace, monospace")[0] == "Consolas"
    assert themes.pick_font("") == ("Segoe UI", None)
    assert themes.pick_font("Zapfino")[0] == "Segoe UI"


def test_radius_is_whole_pixels_within_bounds():
    assert themes.radius_px("12px") == 12 and themes.radius_px("0.75rem") == 12 and themes.radius_px("999px") == themes.MAX_RADIUS
    assert themes.radius_px("-4px") == 0 and themes.radius_px("pill") == 0 and themes.radius_px(None) == 0


def test_contrast_fit_pulls_failing_colours_and_leaves_good_themes_alone():
    good = themes.get_theme("claro")
    assert themes.contrast_issues(good) == [] and themes.fit_contrast(good) == (good, [])
    pale = {"colors": {**good["colors"], "accent": "#ffd0d0", "muted": "#e0e0e0"}, "fonts": good["fonts"]}
    assert {i["fg"] for i in themes.contrast_issues(pale)} == {"accent", "muted"}
    fixed, warnings = themes.fit_contrast(pale)
    assert themes.contrast_issues(fixed) == [] and len(warnings) == 2 and fixed["colors"]["text"] == pale["colors"]["text"]


def test_theme_from_roles_both_modes_pass_the_check():
    for mode in ("light", "dark"):
        theme, warnings = themes.theme_from_roles(sample_roles(), theme_id=f"t-{mode}", name="Aurora", mode=mode)
        assert themes.contrast_issues(theme) == []
        assert theme["colors"]["background"] == sample_roles()[mode]["background"] and theme["colors"]["surface"] == sample_roles()[mode]["surface2"]
        assert theme["fonts"] == {"heading": "Georgia", "body": "Segoe UI"} and theme["radius"] == 12
        assert any("Inter" in w for w in warnings)
    with pytest.raises(CiceroError):
        themes.theme_from_roles(sample_roles(), theme_id="x", name="x", mode="sepia")
    with pytest.raises(CiceroError):
        themes.theme_from_roles({"light": {"background": "red"}}, theme_id="x", name="x")


def test_a_low_contrast_design_system_is_fixed_and_the_fix_is_reported(tmp_path):
    roles = sample_roles()
    roles["light"] = {**roles["light"], "accent": "#ffc0c0"}
    svc = make_services(tmp_path, hub=FakeHub({"ds1": roles}))
    made = svc.theme_from_tokens("ds1", "light")
    assert made["contrast_issues"] == [] and any("accent changed" in w for w in made["warnings"])
    assert made["theme"]["colors"]["accent"] != "#ffc0c0"


def test_theme_from_tokens_saves_a_theme_and_links_it(tmp_path):
    svc = make_services(tmp_path)
    made = svc.theme_from_tokens("ds1", "dark")
    theme = made["theme"]
    assert made["created"] is True and theme["custom"] is True and theme["id"].startswith("ds-ds1-dark-") and theme["name"] == "Aurora (oscuro)"
    assert theme["source_ref"] == "hoard://vitruvius/tokens/ds1" and theme["source_revision"].startswith("sha256:")
    assert [t["id"] for t in svc.themes()][-1] == theme["id"] and len(svc.themes()) == len(themes.THEMES) + 1
    assert svc.hub.calls == [("vitruvius", "tokens_get", {"id": "ds1"})]
    assert svc.hub.links == [(f"hoard://cicero/theme/{theme['id']}", "hoard://vitruvius/tokens/ds1", "derived_from", {"from_label": "Aurora (oscuro)", "to_label": "Aurora"})]
    assert "cicero.theme.created" in svc._emit.types()


def test_theme_from_tokens_is_idempotent_and_a_changed_system_makes_a_new_theme(tmp_path):
    hub = FakeHub()
    svc = make_services(tmp_path, hub=hub)
    first = svc.theme_from_tokens("ds1", "light")
    again = svc.theme_from_tokens("ds1", "light")
    assert again["created"] is False and again["theme"]["id"] == first["theme"]["id"] and len(custom_themes.list_custom(svc)) == 1
    assert len(hub.links) == 1
    other_mode = svc.theme_from_tokens("ds1", "dark")
    assert other_mode["created"] is True
    hub.systems["ds1"] = sample_roles(radius={"md": "2px"})
    changed = svc.theme_from_tokens("ds1", "light")
    assert changed["created"] is True and changed["theme"]["id"] != first["theme"]["id"] and changed["theme"]["radius"] == 2
    assert len(custom_themes.list_custom(svc)) == 3


def test_the_hub_failing_gives_a_clear_error_and_nothing_is_saved(tmp_path):
    svc = make_services(tmp_path, hub=FakeHub(error="hub not reachable at http://127.0.0.1:8810"))
    with pytest.raises(CiceroError) as caught:
        svc.theme_from_tokens("ds1", "light")
    assert "hub not reachable" in str(caught.value) and "linked to the hub" in str(caught.value) and caught.value.code == "family_unavailable"
    assert custom_themes.list_custom(svc) == []
    with pytest.raises(CiceroError):
        make_services(tmp_path / "b", hub=FakeHub()).theme_from_tokens("missing", "light")
    with pytest.raises(CiceroError):  # an older design-system app that sends no roles
        svc2 = make_services(tmp_path / "c")
        svc2._family_call_fn = lambda *a, **k: {"ok": True, "result": {"id": "ds1", "tokens": {}}}
        svc2.theme_from_tokens("ds1", "light")


def test_the_tool_lists_makes_and_applies(tmp_path):
    svc = make_services(tmp_path)
    listed = at.call_tool(svc, "theme_from_tokens", {})
    assert listed["ok"] is True and listed["design_systems"] == [{"id": "ds1", "name": "Aurora"}]
    deck = new_deck(svc)
    made = at.call_tool(svc, "theme_from_tokens", {"tokens_id": "ds1", "mode": "dark", "deck_id": deck["id"]})
    assert made["ok"] is True and made["theme_id"] == made["theme"]["id"] and made["created"] is True and made["applied_to"] == deck["id"] and made["theme"]["radius"] == 12
    assert store.deck_view(svc, deck["id"])["theme"] == made["theme"]["id"]
    # deck_theme sees it, and the contrast report is empty
    assert made["contrast_issues"] == [] and any(t["id"] == made["theme"]["id"] for t in at.call_tool(svc, "deck_theme", {})["themes"])
    with pytest.raises(Exception):
        at.call_tool(svc, "theme_from_tokens", {"tokens_id": "ds1", "mode": "sepia"})


def test_deck_theme_and_update_accept_a_custom_theme_and_still_reject_unknown_ones(tmp_path):
    svc = make_services(tmp_path)
    theme_id = svc.theme_from_tokens("ds1", "light")["theme"]["id"]
    deck = new_deck(svc, theme=theme_id)
    assert deck["theme"] == theme_id and deck["theme_def"]["fonts"]["heading"] == "Georgia"
    assert "theme_def" not in new_deck(svc)
    with pytest.raises(CiceroError):
        store.update_deck(svc, deck["id"], DeckPatch(theme="ds-nope"))
    at.call_tool(svc, "deck_theme", {"deck_id": deck["id"], "theme": "oscuro"})
    assert "theme_def" not in store.deck_view(svc, deck["id"])


def test_a_deck_with_a_custom_theme_renders_checks_and_exports_with_it(tmp_path, deck, services):
    theme_id = services.theme_from_tokens("ds1", "dark")["theme"]["id"]
    store.update_deck(services, deck["id"], {"theme": theme_id})
    store.add_slide(services, deck["id"], title="Dos columnas", blocks=[{"type": "columns", "left_title": "A", "left": ["uno", "dos"], "right_title": "B", "right": ["tres"]}])
    html = services.preview_html(deck["id"])
    assert "--background:#14110f" in html and "--f-h:&quot;Georgia&quot;" in html and "border-radius:12px" in html
    assert isinstance(services.check(deck["id"])["issues"], list)
    exported = services.export(deck["id"], "pptx")
    path, _ = store.export_file(services, exported["id"])
    prs = Presentation(io.BytesIO(path.read_bytes()))
    cards = [s for slide in prs.slides for s in slide.shapes if s.name == "Decoration" and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    assert cards and 0 < cards[0].adjustments[0] <= 0.5
    fonts = {r.font.name for slide in prs.slides for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs}
    assert fonts <= {"Georgia", "Segoe UI"} and "Georgia" in fonts


def test_built_in_themes_render_without_rounding(services, deck):
    store.add_slide(services, deck["id"], title="Dos columnas", blocks=[{"type": "columns", "left_title": "A", "left": ["uno"], "right_title": "B", "right": ["dos"]}])
    html = services.preview_html(deck["id"])
    assert 'class="b" style="' in html and "border-radius:" not in html.split("</style>", 1)[1]


def test_routes(client):
    items = client.get("/api/themes/design-systems").json()["items"]
    assert [d["id"] for d in items] == ["ds1"] and items[0]["roles"]["light"]["background"] == "#fbfaf7"
    made = client.post("/api/themes/from-tokens", json={"tokens_id": "ds1", "mode": "light"})
    assert made.status_code == 201 and made.json()["created"] is True
    assert client.post("/api/themes/from-tokens", json={"tokens_id": "ds1", "mode": "light"}).json()["created"] is False
    assert any(t.get("custom") for t in client.get("/api/themes").json()["items"])
    missing = client.post("/api/themes/from-tokens", json={"tokens_id": "nope"})
    assert missing.status_code == 400 and missing.json()["code"] == "family_unavailable"
    assert client.post("/api/themes/from-tokens", json={"tokens_id": "ds1", "mode": "sepia"}).status_code == 400
    assert client.post("/api/themes/from-tokens", json={"tokens_id": "ds1", "extra": 1}).status_code == 400
