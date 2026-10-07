from concurrent.futures import ThreadPoolExecutor

import pytest
from pptx import Presentation

from cicero_hoard import agent_tools, decks
from cicero_hoard.errors import CiceroError, RevisionConflict
from conftest import png_bytes, new_deck


def test_picture_correction_and_retries_preserve_authored_content(services):
    deck = new_deck(services)
    asset = decks.add_asset(services, deck['id'], png_bytes(), 'figure.png')
    chart = {'type': 'chart', 'chart': 'bar', 'categories': ['A', 'B'],
             'series': [{'name': 'Score', 'values': [2.75, 4.5]}]}
    original = decks.add_slide(services, deck['id'], title='Evidence', layout='image_text',
                               blocks=[chart, {'type': 'bullets', 'items': ['Keep this']}], notes='Original notes')
    args = {'deck_id': deck['id'], 'slide_id': original['id'], 'asset_id': asset['asset_id'], 'caption': 'Original'}
    first = agent_tools.call_tool(services, 'slide_set_image', args)
    with ThreadPoolExecutor(max_workers=4) as pool:
        repeated = list(pool.map(lambda _: agent_tools.call_tool(services, 'slide_set_image', args), range(4)))
    assert first['changed'] and first['block_index'] == 2
    assert all(not r['changed'] and r['revision'] == first['revision'] for r in repeated)
    corrected = agent_tools.call_tool(services, 'slide_set_image', {**args, 'caption': 'Verified'})
    assert corrected['blocks'][:2] == original['blocks']
    assert corrected['title'] == original['title'] and corrected['notes'] == original['notes']
    assert corrected['blocks'][2]['caption'] == 'Verified'
    assert corrected['revision'] == 3
    exported = agent_tools.call_tool(services, 'deck_export', {'deck_id': deck['id'], 'format': 'pptx'})
    slide = Presentation(exported['path']).slides[0]
    assert slide.notes_slide.notes_text_frame.text == 'Original notes'
    assert sum(s.shape_type == 13 for s in slide.shapes) == 1
    assert sum(s.has_chart for s in slide.shapes) == 1
    native_chart = next(s.chart for s in slide.shapes if s.has_chart)
    assert list(native_chart.series[0].values) == [2.75, 4.5]
    reverted = decks.revert_slide(services, deck['id'], original['id'], 1)
    assert reverted['blocks'] == original['blocks'] and reverted['notes'] == original['notes']


def test_multiple_images_caption_semantics_and_invalid_edits_are_atomic(services):
    deck = new_deck(services)
    asset = decks.add_asset(services, deck['id'], png_bytes(), 'image.png')
    slide = decks.add_slide(services, deck['id'], blocks=[{'type': 'text', 'text': 'Body'}])
    args = {'deck_id': deck['id'], 'slide_id': slide['id'], 'asset_id': asset['asset_id']}
    agent_tools.call_tool(services, 'slide_set_image', {**args, 'caption': 'First'})
    second = agent_tools.call_tool(services, 'slide_set_image', {**args, 'image_index': 1, 'caption': 'Second'})
    same = agent_tools.call_tool(services, 'slide_set_image', args)
    assert not same['changed'] and same['blocks'][1]['caption'] == 'First'
    cleared = agent_tools.call_tool(services, 'slide_set_image', {**args, 'image_index': 1, 'caption': None})
    assert cleared['blocks'][1] == second['blocks'][1] and 'caption' not in cleared['blocks'][2]
    for bad in ({'image_index': 3}, {'asset_id': 'missing'}, {'expected_revision': 1}):
        with pytest.raises(CiceroError):
            agent_tools.call_tool(services, 'slide_set_image', {**args, **bad})
        assert decks.get_slide(services, deck['id'], slide['id']) == {k: v for k, v in cleared.items() if k not in ('changed', 'block_index')}
    other = new_deck(services)
    foreign = decks.add_asset(services, other['id'], png_bytes(), 'foreign.png')
    with pytest.raises(CiceroError):
        agent_tools.call_tool(services, 'slide_set_image', {**args, 'asset_id': foreign['asset_id']})


def test_rest_image_edit_preserves_omitted_caption_and_reports_conflict(client, tmp_path):
    deck = client.post('/api/decks', json={'title': 'API figure'}).json()
    slide = client.post(f"/api/decks/{deck['id']}/slides", json={'title': 'Keep', 'notes': 'Notes'}).json()
    file = tmp_path/'image.png'; file.write_bytes(png_bytes())
    asset = client.post(f"/api/decks/{deck['id']}/assets/import", json={'path': str(file)}).json()
    url = f"/api/decks/{deck['id']}/slides/{slide['id']}/image"
    first = client.put(url, json={'asset_id': asset['asset_id'], 'caption': 'Keep caption'}).json()
    same = client.put(url, json={'asset_id': asset['asset_id']}).json()
    assert same['revision'] == first['revision'] and not same['changed']
    assert same['blocks'][0]['caption'] == 'Keep caption' and same['notes'] == 'Notes'
    conflict = client.put(url, json={'asset_id': asset['asset_id'], 'expected_revision': 1})
    assert conflict.status_code == 409 and conflict.json()['code'] == 'conflict'
    assert client.put(url, json={'asset_id': asset['asset_id'], 'caption': None}).json()['blocks'] == [{'type': 'image', 'asset_id': asset['asset_id']}]
