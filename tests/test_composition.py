import io
import pytest
from pydantic import ValidationError
from pptx import Presentation
from PIL import Image
from cicero_hoard.models import ImageBlock,TextBlock,TimelineBlock
from cicero_hoard.layouts import build_plan
from cicero_hoard.themes import theme_of
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.export_md import build_markdown
from cicero_hoard.render import render_document
from cicero_hoard import native_slides

def deck_with(blocks,layout='freeform'):
 return {'title':'Reusable composition','language':'en','theme':'claro','slides':[{'id':'s','position':1,'title':'Composition','layout':layout,'blocks':blocks,'notes':'Notes'}]}

@pytest.mark.parametrize('frame',[[.8,0,.4,1],[0,0,0,1],[0,-.1,1,1],[0,0,float('nan'),1]])
def test_frames_are_finite_inside_the_content_area(frame):
 with pytest.raises(ValidationError):TextBlock(frame=frame)

def test_free_composition_preserves_ids_fonts_and_reports_overlaps():
 deck=deck_with([{'type':'text','id':'label','text':'Left','frame':[0,0,.4,.5],'font_size':48},
                  {'type':'text','text':'Right','frame':[.5,0,.5,.5],'font_size':24}])
 plan=build_plan(deck['slides'][0],deck_title=deck['title'],theme=theme_of(deck),number=1,total=1)
 texts=[b for b in plan.boxes if b.kind=='text' and b.role not in ('title','footer')]
 assert texts[0].x==72 and texts[0].w==pytest.approx(454.4) and texts[0].paras[0].size==48
 assert texts[0].object_id=='label-0' and texts[1].x>texts[0].x+texts[0].w and not plan.overflow
 html,_=render_document(deck,assets=lambda _:None)
 assert 'data-id="label-0"' in html
 deck['slides'][0]['blocks'][1]['frame']=[0,0,.5,.5]
 plan=build_plan(deck['slides'][0],deck_title=deck['title'],theme=theme_of(deck),number=1,total=1)
 assert any('overlapping' in warning for warning in plan.overflow_where)

def test_cover_crop_matches_web_and_native_focus(tmp_path):
 path=tmp_path/'wide.png';Image.new('RGB',(400,200),'#31a4c6').save(path)
 assets=lambda _: {'path':path,'mime':'image/png','width':400,'height':200}
 block=ImageBlock(asset_id='a',fit='cover',position=[1,.5],frame=[0,0,.25,1]).model_dump(exclude_none=True)
 deck=deck_with([block]);prs=Presentation(io.BytesIO(build_pptx(deck,assets=assets)))
 pic=next(s for s in prs.slides[0].shapes if s.shape_type==13)
 assert pic.crop_left>.5 and pic.crop_right==0 and pic.crop_top==0
 html,_=render_document(deck,assets=assets)
 assert 'object-fit:cover' in html and 'object-position:100.0%' in html
 with pytest.raises(ValidationError):ImageBlock(asset_id='a',position=[1.1,.5])

def timeline_deck():
 block=TimelineBlock(effect='stages',items=[{'date':'Phase A','label':'Collect','detail':'Source records'},{'date':'Phase B','label':'Compare','detail':'Candidate methods'},{'date':'Phase C','label':'Deliver','detail':'Reviewed results','emphasis':True}]).model_dump(exclude_none=True)
 return deck_with([block],layout='evidence')

def test_timeline_has_editable_objects_click_groups_and_text_exports():
 deck=timeline_deck();plan=build_plan(deck['slides'][0],deck_title=deck['title'],theme=theme_of(deck),number=1,total=1)
 assert {b.build for b in plan.boxes if b.build}=={1,2,3} and sum(b.kind=='marker' for b in plan.boxes)==3 and not plan.overflow
 prs=Presentation(io.BytesIO(build_pptx(deck,assets=lambda _:None)))
 text=' '.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
 assert all(label in text for label in ('Collect','Compare','Deliver'))
 assert 'Phase A' in build_markdown(deck).decode()
 html,_=render_document(deck,assets=lambda _:None)
 assert 'data-fragment-index="2"' in html

@pytest.mark.skipif(not native_slides.available(),reason='Optional local native slide engine')
def test_timeline_opens_and_animates_in_native_engine(tmp_path):
 raw=build_pptx(timeline_deck(),assets=lambda _:None)
 report=native_slides.render(raw,tmp_path/'timeline-native')
 assert report['slide_count']==1 and report['animation_count']>=3
 assert (tmp_path/'timeline-native'/'slide-1.png').is_file()
