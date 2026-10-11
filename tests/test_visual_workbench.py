import io
import json
import zipfile
from copy import deepcopy
import pytest
from pydantic import ValidationError
from pptx import Presentation
from cicero_hoard.models import ChartBlock,DiagramBlock
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.render import chart_svg,render_document,present_document
from cicero_hoard.themes import get_theme
from cicero_hoard.layouts import build_plan
from cicero_hoard.agent_tools import call_tool
from cicero_hoard import decks
from cicero_hoard.errors import RevisionConflict
from conftest import new_deck

KINDS=['bar','bar_horizontal','line','area','pie','donut','scatter','bubble','dot','dumbbell','band','heatmap','waterfall','funnel']

def test_band_upper_bound_stays_in_the_packaged_chart_canvas():
 import xml.etree.ElementTree as ET
 block={'chart':'band','categories':['A','B'],'series':[{'name':'Range','values':[2,3],'x_values':[1,2],'lower':[1,1],'upper':[100,100]}]}
 svg=chart_svg(block,1000,500,get_theme('claro'),'en','Arial')
 root=ET.fromstring(svg)
 polygons=root.findall('.//{http://www.w3.org/2000/svg}polygon')
 assert polygons
 for polygon in polygons:
  points=[float(value) for value in polygon.attrib['points'].replace(',', ' ').split()]
  assert len(points)>=6 and len(points)%2==0
  assert all(0<=y<=500 for y in points[1::2])

def chart(kind):
    b={'type':'chart','chart':kind,'categories':['First','Second','Third'],'unit':'kPa',
       'series':[{'name':'Measured','values':[3.49,4.26,6.03]}]}
    if kind=='dumbbell':b['series'].append({'name':'Comparison','values':[4.26,3.49,7.83]})
    if kind in ('scatter','bubble','band'):b['series'][0]['x_values']=[2,10,90]
    if kind=='bubble':b['series'][0]['sizes']=[1,4,9]
    if kind=='band':b['series'][0].update(lower=[2,3,5],upper=[4,5,7])
    return ChartBlock.model_validate(b).model_dump(exclude_none=True)

@pytest.mark.parametrize('kind',KINDS)
def test_every_chart_renders_and_exports_editable_evidence(kind):
    b=chart(kind)
    svg=chart_svg(b,1000,460,get_theme('claro'),'es','Arial')
    assert '<svg' in svg and len(svg)>500
    deck={'title':'Generic data comparison','language':'es','theme':'claro','slides':[{'id':'s','position':1,'title':'Evidence','layout':'evidence','blocks':[b],'notes':'Measured data.'}]}
    raw=build_pptx(deck,assets=lambda _:None)
    prs=Presentation(io.BytesIO(raw))
    assert any(s.has_chart or s.has_table or s.shape_type==1 for s in prs.slides[0].shapes)
    if kind not in ('heatmap','funnel'):
        charts=[s.chart for s in prs.slides[0].shapes if s.has_chart]
        assert len(charts)==1
        with zipfile.ZipFile(io.BytesIO(raw)) as z:assert any(n.startswith('ppt/embeddings/') for n in z.namelist())

def test_numeric_coordinates_and_interval_contracts():
    with pytest.raises(ValidationError):ChartBlock.model_validate({**chart('band'),'series':[{'name':'Invalid','values':[3],'x_values':[2],'lower':[4],'upper':[5]}]})
    with pytest.raises(ValidationError):ChartBlock.model_validate({**chart('scatter'),'series':[{'name':'Invalid','values':[1,2],'x_values':[2]}]})
    with pytest.raises(ValidationError):ChartBlock.model_validate({**chart('bar'),'annotations':[{'text':'Outside','category':7}]})
    with pytest.raises(ValidationError):DiagramBlock(nodes=[{'id':'a','label':'A'}],edges=[{'source':'a','target':'absent'}])
    with pytest.raises(ValidationError):DiagramBlock(nodes=[{'id':'a','label':'A'},{'id':'a','label':'Duplicate'}])

@pytest.mark.parametrize('blocks',[
    [{'type':'metrics','items':[{'value':'120','label':'Eligible'}]},{'type':'text','text':'Takeaway'}],
    [{'type':'text','text':'Takeaway'},{'type':'bullets','items':['Retained content']}],
])
def test_evidence_never_drops_non_chart_content(blocks):
    plan=build_plan({'layout':'evidence','title':'Proof','blocks':blocks},deck_title='Proof',theme=get_theme('claro'),number=1,total=1,lang='en',assets=lambda _:None)
    text=' '.join(p.text for b in plan.boxes for p in b.paras)
    assert 'Takeaway' in text
    if blocks[0]['type']=='metrics':assert '120' in text and 'Eligible' in text
    else:assert 'Retained content' in text

def test_targeted_visual_edit_is_idempotent_and_preserves_other_blocks(services):
    deck=new_deck(services);slide=decks.add_slide(services,deck['id'],title='Generic',notes='Retain these notes',blocks=[{'type':'text','text':'Retained'}])
    args={'deck_id':deck['id'],'slide_id':slide['id'],'block_index':1,'block':chart('bar_horizontal'),'expected_revision':1}
    out=call_tool(services,'slide_visual_update',args)
    assert out['notes']=='Retain these notes' and out['blocks'][0]['text']=='Retained' and out['revision']==2
    args['expected_revision']=2
    again=call_tool(services,'slide_visual_update',args)
    assert again['changed'] is False and again['revision']==2
    args['expected_revision']=1
    with pytest.raises(RevisionConflict):call_tool(services,'slide_visual_update',args)

def test_agent_discovery_variants_scene_and_live_link(services):
    deck=new_deck(services);slide=decks.add_slide(services,deck['id'],title='Data',blocks=[chart('bar_horizontal')],notes='Notes')
    catalog=call_tool(services,'visual_catalog',{})
    assert len(catalog['charts'])==14
    schema=call_tool(services,'visual_catalog',{'kind':'charts'})
    assert 'x_values' in json.dumps(schema) and 'annotations' in json.dumps(schema)
    variants=call_tool(services,'slide_variants',{'deck_id':deck['id'],'slide_id':slide['id']})
    assert len(variants['variants'])==3 and all(v['content_changed'] is False for v in variants['variants'])
    assert decks.get_slide(services,deck['id'],slide['id'])['revision']==1
    scene=call_tool(services,'deck_scene',{'deck_id':deck['id']})
    assert scene['canvas']==[1280,720] and scene['slides'][0]['objects']
    link=call_tool(services,'deck_present',{'deck_id':deck['id']})
    assert link['url'].endswith('?mode=present')

def test_diagram_objects_and_click_groups_remain_native():
    b={'type':'diagram','effect':'stages','nodes':[{'id':'a','label':'Input'},{'id':'b','label':'Transform'},{'id':'c','label':'Result'}],
       'edges':[{'source':'a','target':'b'},{'source':'b','target':'c'}]}
    deck={'title':'Workflow','language':'es','theme':'claro','slides':[{'id':'s','position':1,'layout':'diagram','title':'Workflow','blocks':[b],'notes':'Source process'}]}
    html,_=render_document(deck,assets=lambda _:None)
    assert 'data-fragment-index="2"' in html
    live=present_document(html)
    assert 'new Reveal' in live and 'prefers-reduced-motion' in live and 'Source process' not in live
    prs=Presentation(io.BytesIO(build_pptx(deck,assets=lambda _:None)))
    names=[shape.name for shape in prs.slides[0].shapes]
    assert sum(name.startswith('Cicero-step:') for name in names)>=6
    assert not any(shape.shape_type==13 for shape in prs.slides[0].shapes)

def test_motion_does_not_remove_static_export_when_native_engine_is_missing(services):
    from cicero_hoard.errors import PdfUnavailable
    deck=new_deck(services)
    decks.add_slide(services,deck['id'],title='A staged explanation',blocks=[{'type':'text','text':'Retained','build':1,'effect':'fade'}])
    def unavailable(*args):raise PdfUnavailable('Unavailable')
    services._native_render_fn=unavailable
    out=services.export(deck['id'],'pptx')
    path,_=decks.export_file(services,out['id'])
    assert out['warnings'] and Presentation(str(path)).slides[0].shapes

def test_variant_previews_are_readonly(client):
    deck=client.post('/api/decks',json={'title':'Composition review'}).json()
    slide=client.post(f'/api/decks/{deck["id"]}/slides',json={'title':'Retained title','notes':'Retained notes','blocks':[chart('bar_horizontal')]}).json()
    url=f'/api/decks/{deck["id"]}/slides/{slide["id"]}'
    variants=client.get(url+'/variants').json()
    assert all('variant.html?layout=' in v['preview_url'] for v in variants['variants'])
    preview=client.get(url+'/variant.html?layout=evidence')
    assert preview.status_code==200 and 'Retained title' in preview.text
    current=client.get(url).json()
    assert current['revision']==slide['revision'] and current['notes']=='Retained notes' and current['layout']==slide['layout']
    assert client.get(url+'/variant.html?layout=absent').status_code==400

def test_visual_runtime_is_packaged():
    from pathlib import Path
    root=Path(__file__).resolve().parent.parent/'cicero_hoard/static/visual-runtime'
    assert all((root/name).is_file() for name in ['engine.mjs','reveal.js','reveal.css','LICENSES.txt'])
