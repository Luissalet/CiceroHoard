"""Exports: PPTX opened again with python-pptx, standalone HTML, Markdown, PDF and its fallback error."""
from __future__ import annotations

import io
import re

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from cicero_hoard import decks, export_md, export_pdf
from cicero_hoard.errors import CiceroError, PdfUnavailable
from cicero_hoard.export_html import build_html
from cicero_hoard.export_pptx import build_pptx
from conftest import make_services, new_deck, png_bytes
from sample_deck import sample_deck

NO_ASSETS = lambda _id: None  # noqa: E731


def open_pptx(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def texts(slide) -> list[str]:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
    return out


# ---------------- PPTX ----------------

def test_pptx_structure():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    assert len(prs.slides) == 11
    assert prs.slide_width == Emu(12192000) and prs.slide_height == Emu(6858000)  # 16:9, 13.333 x 7.5 in
    assert prs.core_properties.title == "Revisión <trimestral> & \"cifras\"" and prs.core_properties.language == "es-ES"


def test_pptx_text_is_real_text_and_editable():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    s3 = prs.slides[2]
    assert s3.shapes.title is not None and s3.shapes.title.text_frame.text == "Las ventas crecieron un 12%"
    body = [t for t in texts(s3) if "1.250.000" in t]
    assert body and body[0].count("\n") == 5  # six bullet paragraphs
    assert not [sh for sh in s3.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]  # no text baked into pictures


def test_pptx_bullets_are_paragraph_bullets():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    shape = next(sh for sh in prs.slides[2].shapes if sh.has_text_frame and "1.250.000" in sh.text_frame.text)
    paras = shape.text_frame.paragraphs
    assert len(paras) == 6
    for p in paras:
        ppr = p._p.pPr
        assert ppr is not None and ppr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None


def test_pptx_notes():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Notas del orador número 3."


def test_pptx_native_charts():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    kinds = []
    for i in (5, 6, 7):
        charts = [sh.chart for sh in prs.slides[i].shapes if sh.has_chart]
        assert len(charts) == 1
        kinds.append(str(charts[0].chart_type))
    assert "COLUMN" in kinds[0] and "LINE" in kinds[1] and "PIE" in kinds[2]
    bar = next(sh.chart for sh in prs.slides[5].shapes if sh.has_chart)
    assert list(bar.plots[0].categories) == ["Hogar", "Moda", "Deporte"]
    assert [list(s.values) for s in bar.plots[0].series] == [[380, 310, 150], [475, 340, 160]]


def test_pptx_font_downscaling_matches_the_plan():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))

    def sizes(slide):
        out = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                for p in sh.text_frame.paragraphs:
                    out += [r.font.size.pt for r in p.runs if r.font.size]
        return out

    long_title = prs.slides[3].shapes.title
    short_title = prs.slides[2].shapes.title
    assert min(r.font.size.pt for p in long_title.text_frame.paragraphs for r in p.runs) < min(r.font.size.pt for p in short_title.text_frame.paragraphs for r in p.runs)
    assert max(sizes(prs.slides[2])) >= 12  # body text never goes below 16 px (12 pt); only the footer is smaller


def test_pptx_picture_and_webp_conversion(tmp_path):
    for fmt, mime, ext in (("PNG", "image/png", "png"), ("WEBP", "image/webp", "webp")):
        path = tmp_path / f"a1.{ext}"
        path.write_bytes(png_bytes(fmt=fmt))
        lookup = lambda aid, path=path, mime=mime: ({"id": "a1", "mime": mime, "path": path, "width": 64, "height": 48} if aid == "a1" else None)  # noqa: E731
        prs = open_pptx(build_pptx(sample_deck(image_asset="a1"), assets=lookup))
        pics = [sh for sh in prs.slides[9].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1 and pics[0].image.content_type in ("image/png", "image/jpeg")


def test_pptx_missing_image_does_not_break():
    prs = open_pptx(build_pptx(sample_deck(image_asset="ghost"), assets=NO_ASSETS))
    assert len(prs.slides) == 11


def test_pptx_english_language_and_all_themes():
    for theme in ("claro", "oscuro", "editorial", "carmesi", "pizarra", "tecnico", "oceano"):
        prs = open_pptx(build_pptx(sample_deck(theme, language="en"), assets=NO_ASSETS))
        assert len(prs.slides) == 11 and prs.core_properties.language == "en-US"


def test_pptx_text_stays_inside_the_slide():
    prs = open_pptx(build_pptx(sample_deck(), assets=NO_ASSETS))
    for slide in prs.slides:
        for sh in slide.shapes:
            assert sh.left >= -10000 and sh.top >= -10000
            assert sh.left + sh.width <= prs.slide_width + 10000 and sh.top + sh.height <= prs.slide_height + 10000


# ---------------- HTML ----------------

def test_html_is_standalone_and_leaves_notes_out(tmp_path):
    path = tmp_path / "a1.png"
    path.write_bytes(png_bytes())
    lookup = lambda aid: ({"id": "a1", "mime": "image/png", "path": path, "width": 64, "height": 48} if aid == "a1" else None)  # noqa: E731
    html = build_html(sample_deck(image_asset="a1"), assets=lookup).decode("utf-8")
    assert html.startswith("<!doctype html>") and "data:image/png;base64," in html
    assert "Notas del orador" not in html
    assert not re.search(r"(src|href)\s*=\s*[\"']https?://", html) and "/api/" not in html
    assert html.count("<section") == 11 and 'data-mode="present"' in html


# ---------------- Markdown ----------------

def test_markdown_export():
    md = export_md.build_markdown(sample_deck()).decode("utf-8")
    assert md.count("\n---\n") == 10
    assert "# Las ventas crecieron un 12%" in md and "- La rotación subió al 9%" in md
    assert "<!-- Notas" in md or "<!-- Notes" in md
    assert "> Encuentro lo que busco" in md and "— Cliente de ejemplo" in md
    assert "| Hogar |" in md  # charts as tables
    assert "**Antes**" in md


def test_markdown_notes_cannot_break_out_of_the_comment():
    deck = sample_deck()
    deck["slides"][2]["notes"] = "cierre --> <script>x</script> --"
    md = export_md.build_markdown(deck).decode("utf-8")
    comment = re.search(r"<!--(.*?)-->", md.split("---")[2], flags=re.S).group(1)
    assert "-->" not in comment and "--" not in comment


# ---------------- PDF ----------------

def test_pdf_unavailable_is_a_clear_error(monkeypatch):
    monkeypatch.setattr(export_pdf, "_launch", lambda p: (_ for _ in ()).throw(PdfUnavailable("No Chromium is available")))
    with pytest.raises(PdfUnavailable) as info:
        export_pdf.build_pdf(sample_deck(), assets=NO_ASSETS)
    assert info.value.code == "pdf_unavailable" and "playwright install chromium" in export_pdf.INSTALL_HINT


def test_pdf_generic_browser_failure_is_pdf_unavailable(monkeypatch):
    def boom(p):
        raise RuntimeError("browser crashed\nstack")
    monkeypatch.setattr(export_pdf, "_launch", boom)
    with pytest.raises(PdfUnavailable) as info:
        export_pdf.build_pdf(sample_deck(), assets=NO_ASSETS)
    assert "browser crashed" in str(info.value) and "stack" not in str(info.value)


@pytest.mark.skipif(not export_pdf.pdf_available(), reason="no Chromium for the PDF export here")
def test_pdf_real_render():
    data = export_pdf.build_pdf(sample_deck(), assets=NO_ASSETS)
    assert data.startswith(b"%PDF-")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) == 11
    box = reader.pages[0].mediabox
    assert abs(float(box.width) / float(box.height) - 16 / 9) < 0.02
    assert "Revisión trimestral" in reader.pages[0].extract_text() or "trimestral" in reader.pages[0].extract_text()


# ---------------- through the service ----------------

def test_service_export_all_formats_and_records(services, deck):
    got = {}
    for fmt in ("pptx", "html", "md"):
        exp = services.export(deck["id"], fmt)
        got[fmt] = exp
        assert exp["format"] == fmt and exp["bytes"] > 100 and exp["filename"].endswith(f".{fmt}")
        path, _ = decks.export_file(services, exp["id"])
        assert path.is_file() and path.stat().st_size == exp["bytes"]
    assert len(open_pptx(decks.export_file(services, got["pptx"]["id"])[0].read_bytes()).slides) == len(deck["slides"])
    assert len(decks.list_exports(services, deck["id"])["items"]) == 3
    assert "cicero.deck.exported" in services._emit.types()


def test_service_export_rules(services):
    d = new_deck(services)
    with pytest.raises(CiceroError) as info:
        services.export(d["id"], "pptx")
    assert info.value.code == "no_slides"
    decks.add_slide(services, d["id"], title="x")
    with pytest.raises(CiceroError):
        services.export(d["id"], "odp")


def test_service_pdf_uses_injected_function_and_reports_failure(tmp_path):
    svc = make_services(tmp_path, pdf_fn=lambda deck, assets: b"%PDF-1.4 fake")
    d = new_deck(svc)
    decks.add_slide(svc, d["id"], title="x")
    assert svc.export(d["id"], "pdf")["bytes"] == len(b"%PDF-1.4 fake")

    def fail(deck, assets):
        raise PdfUnavailable("nope")
    svc2 = make_services(tmp_path / "b", pdf_fn=fail)
    d2 = new_deck(svc2)
    decks.add_slide(svc2, d2["id"], title="x")
    with pytest.raises(PdfUnavailable):
        svc2.export(d2["id"], "pdf")
