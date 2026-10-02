"""Text extraction from documents and the source endpoints of the store."""
from __future__ import annotations

import io
import zipfile

import pytest

from pathlib import Path

from cicero_hoard import decks, extract
from cicero_hoard.errors import CiceroError, Refused
from conftest import minimal_pdf, new_deck


def _docx(paragraphs, heading=None) -> bytes:
    from docx import Document

    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Valor 42"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Resultados"
    slide.placeholders[1].text = "Primer punto\nSegundo punto"
    slide.notes_slide.notes_text_frame.text = "Nota del orador"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_text_and_markdown():
    got = extract.extract_bytes("Hola\r\n\r\n\r\n\r\nmundo  \n".encode(), "nota.txt")
    assert got.kind == "text" and got.text == "Hola\n\nmundo" and got.title == "nota"
    md = extract.extract_bytes(b"# Titulo\n\ntexto", "a.md")
    assert md.kind == "markdown" and md.text.startswith("# Titulo")


def test_encodings():
    assert extract.decode_text("Año".encode("utf-8-sig")) == "Año"
    assert extract.decode_text("Año".encode("cp1252")) == "Año"
    assert extract.decode_text("Año".encode("utf-16")) == "Año"
    with pytest.raises(CiceroError):
        extract.decode_text(b"abc\x00\x01\x02")


def test_docx_text_and_tables():
    got = extract.extract_bytes(_docx(["Primer párrafo con 1.500 euros."], heading="Informe"), "informe.docx")
    assert got.kind == "docx" and "Primer párrafo" in got.text and "Valor 42" in got.text


def test_pptx_text_and_notes():
    got = extract.extract_bytes(_pptx(), "deck.pptx")
    assert got.kind == "pptx" and "Resultados" in got.text and "Segundo punto" in got.text


def test_pdf_text():
    got = extract.extract_bytes(minimal_pdf("Ventas de septiembre 1250"), "informe.pdf")
    assert got.kind == "pdf" and "Ventas de septiembre" in got.text


@pytest.mark.parametrize("data,name", [
    (b"", "a.txt"),                       # empty
    (b"MZ\x00\x00binary", "a.txt"),       # binary
    (b"hola", "a.exe"),                   # unsupported
    (b"hola", "sin_extension"),
    (b"not a pdf", "a.pdf"),
    (b"not a zip", "a.docx"),
    (b"PKnotreallyzip", "a.pptx"),
    (b"   \n  ", "a.txt"),                # no readable text
])
def test_refused_inputs(data, name):
    with pytest.raises(CiceroError):
        extract.extract_bytes(data, name)


def test_zip_without_expected_part_is_refused():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(CiceroError):
        extract.extract_bytes(buf.getvalue(), "fake.docx")


def test_size_limit_and_truncation():
    with pytest.raises(CiceroError):
        extract.extract_bytes(b"a" * 2000, "big.txt", max_bytes=1000)
    long = ("palabra " * 60000).encode()
    got = extract.extract_bytes(long, "largo.txt")
    assert got.truncated and len(got.text) == extract.MAX_TEXT_CHARS


def test_store_sources_lifecycle(services):
    d = new_deck(services)
    s = decks.add_source_text(services, d["id"], "Uno", "Texto de la fuente con 3 cifras: 10, 20 y 30.")
    assert s["id"] == 1 and s["kind"] == "text" and s["chars"] > 10 and "text" not in s
    f = decks.add_source_file(services, d["id"], _docx(["Contenido docx"]), "doc.docx")
    assert f["kind"] == "docx" and f["title"] == "doc"
    assert [i["id"] for i in decks.list_sources(services, d["id"])["items"]] == [1, 2]
    part = decks.get_source(services, d["id"], 1, offset=6, max_chars=8)
    assert part["text"] == "de la fu" and part["offset"] == 6
    with pytest.raises(CiceroError):
        decks.add_source_text(services, d["id"], "vacío", "   ")
    decks.remove_source(services, d["id"], 1)
    assert [i["id"] for i in decks.list_sources(services, d["id"])["items"]] == [2]


def test_source_path_rules(services, tmp_path):
    d = new_deck(services)
    ok = tmp_path / "ok.md"
    ok.write_text("# Hola\n\ncontenido", encoding="utf-8")
    assert decks.add_source_path(services, d["id"], str(ok))["kind"] == "markdown"
    bad = tmp_path / "x.exe"
    bad.write_bytes(b"MZ")
    with pytest.raises(Refused):
        decks.add_source_path(services, d["id"], str(bad))
    with pytest.raises(CiceroError):
        decks.add_source_path(services, d["id"], str(tmp_path / "missing.txt"))
    with pytest.raises(CiceroError):
        decks.add_source_path(services, d["id"], str(tmp_path))


def test_source_path_respects_file_roots(tmp_path):
    from conftest import make_services

    inside = tmp_path / "docs"
    inside.mkdir()
    (inside / "a.txt").write_text("dentro", encoding="utf-8")
    outside = tmp_path / "b.txt"
    outside.write_text("fuera", encoding="utf-8")
    svc = make_services(tmp_path, file_roots=(inside,))
    d = new_deck(svc)
    assert decks.add_source_path(svc, d["id"], str(inside / "a.txt"))["chars"] == 6
    with pytest.raises(Refused):
        decks.add_source_path(svc, d["id"], str(outside))
    with pytest.raises(Refused):
        decks.add_source_path(svc, d["id"], str(inside / ".." / "b.txt"))


# ---------------- the shared readers (Hoard Link) ----------------

def _blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_docx_keeps_headings_lists_and_tables():
    from docx import Document

    doc = Document()
    doc.add_heading("Informe anual", level=1)
    doc.add_paragraph("Texto inicial.")
    doc.add_heading("Ventas", level=2)
    doc.add_paragraph("Primer punto", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Valor 42"
    buf = io.BytesIO()
    doc.save(buf)
    got = extract.extract_bytes(buf.getvalue(), "informe.docx")
    assert got.kind == "docx"
    assert "## Informe anual" in got.text and "## Ventas" in got.text and "- Primer punto" in got.text and "Concepto | Valor 42" in got.text


def test_a_text_file_named_docx_is_refused_not_read_as_text():
    with pytest.raises(CiceroError):
        extract.extract_bytes("hola, esto es texto".encode(), "falso.docx")
    # a pptx renamed .docx is not a Word document either
    with pytest.raises(CiceroError):
        extract.extract_bytes(_pptx(), "falso.docx")


def test_a_zip_bomb_is_refused_by_the_shared_guard():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w:document/>")
        z.writestr("bomb.bin", b"\x00" * (60 * 1024 * 1024))
    with pytest.raises(CiceroError) as info:
        extract.extract_bytes(buf.getvalue(), "bomba.docx")
    assert "refused" in str(info.value)


def test_a_scanned_pdf_goes_to_the_family_ocr_when_it_runs(monkeypatch):
    calls = []

    def fake_extract(path, **kw):
        calls.append((Path(path).suffix, kw))
        return {"ok": True, "via": "kafka", "text": "Texto reconocido por OCR 2024", "needs_ocr": False}

    monkeypatch.setattr(extract.fam_docs, "extract", fake_extract)
    got = extract.extract_bytes(_blank_pdf(), "escaneo.pdf")
    assert got.kind == "pdf" and "reconocido por OCR" in got.text
    assert calls == [(".pdf", {"ocr": "auto", "local_fallback": False})]


def test_a_scanned_pdf_without_the_ocr_service_says_it_needs_ocr(monkeypatch):
    monkeypatch.setattr(extract.fam_docs, "extract", lambda path, **kw: {"ok": False, "error": "hub unreachable"})
    with pytest.raises(CiceroError) as info:
        extract.extract_bytes(_blank_pdf(), "escaneo.pdf")
    assert "OCR" in str(info.value)
