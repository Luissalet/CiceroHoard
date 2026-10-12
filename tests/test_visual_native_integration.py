import pytest
import io
from pptx import Presentation
from cicero_hoard import native_slides
from cicero_hoard.export_pptx import build_pptx
from test_visual_workbench import KINDS,chart

@pytest.mark.parametrize('kind',['bar','bar_horizontal','line','donut'])
def test_native_defaults_show_exact_sparse_values_without_forced_decimals(kind):
    b={'type':'chart','chart':kind,'categories':['Carga','Proceso','Descarga'],'unit':'s',
       'series':[{'name':'Observed','values':[24,38,18]}]}
    raw=build_pptx({'title':'Defaults','language':'en','theme':'claro','slides':[
       {'id':'1','position':1,'title':'Values','layout':'evidence','blocks':[b]}]},assets=lambda _:None)
    native=next(s.chart for s in Presentation(io.BytesIO(raw)).slides[0].shapes if s.has_chart)
    assert native.plots[0].has_data_labels and native.plots[0].data_labels.show_value
    assert native.plots[0].data_labels.number_format=='#,##0'
    if kind!='donut':assert native.value_axis.axis_title.text_frame.text=='s'

def test_native_explicit_label_hide_preserved_and_fractional_format():
    b={'type':'chart','chart':'bar','categories':['A','B'],'unit':'F1','show_values':False,
       'series':[{'name':'Observed','values':[.61,.785]}]}
    raw=build_pptx({'title':'Defaults','language':'en','theme':'claro','slides':[
       {'id':'1','position':1,'title':'Values','layout':'evidence','blocks':[b]}]},assets=lambda _:None)
    native=next(s.chart for s in Presentation(io.BytesIO(raw)).slides[0].shapes if s.has_chart)
    assert not native.plots[0].has_data_labels
    assert native.value_axis.tick_labels.number_format=='#,##0.###'

@pytest.mark.parametrize('kind',['bar','pie','donut','funnel'])
def test_native_all_axis_captions_preserved_and_separated(kind):
    b={'type':'chart','chart':kind,'categories':['A','B'],'unit':'mg','y_label':'Concentration','x_label':'Period',
       'series':[{'name':'Observed','values':[12,20]}]}
    raw=build_pptx({'title':'Captions','language':'en','theme':'claro','slides':[
       {'id':'1','position':1,'title':'Values','layout':'evidence','blocks':[b]}]},assets=lambda _:None)
    shapes=Presentation(io.BytesIO(raw)).slides[0].shapes
    text=[s for s in shapes if s.has_text_frame and s.text in ('Concentration','mg','Period')]
    assert set(s.text for s in text)==({'mg','Period'} if kind=='bar' else {'Concentration','mg','Period'})
    assert len(set(s.top for s in text))==len(text)
    if kind=='bar':assert next(s.chart for s in shapes if s.has_chart).value_axis.axis_title.text_frame.text=='Concentration'

@pytest.mark.skipif(not native_slides.available(),reason='Optional local native slide engine')
def test_all_visual_types_open_together_with_native_labels(tmp_path):
    slides=[]
    for i,kind in enumerate(KINDS,1):
        b=chart(kind);b['series'][0]['values']=[3,4,6];b['show_values']=True
        if kind=='band':b['series'][0].update(lower=[2,3,5],upper=[4,5,7])
        slides.append({'id':str(i),'position':i,'title':kind,'layout':'evidence',
            'blocks':[b,{'type':'text','text':'Fictional demonstration data.'}]})
    raw=build_pptx({'title':'Generic visuals','language':'en','theme':'claro','slides':slides},assets=lambda _:None)
    report=native_slides.render(raw,tmp_path/'native')
    assert report['slide_count']==14 and all((tmp_path/'native'/f'slide-{n}.png').is_file() for n in range(1,15))
