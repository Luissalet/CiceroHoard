import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from cicero_hoard import agent_tools, chrome, decks
from cicero_hoard.errors import CiceroError
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.layouts import build_plan
from cicero_hoard.models import DeckCreate
from cicero_hoard.render import render_document
from cicero_hoard.themes import theme_of
from sample_deck import sample_deck
from test_templates import template

NO_ASSETS = lambda _id: None


def test_labels_keep_history_and_exclude_the_appendix():
    assert chrome.label(None, 3, 9) == '3 / 9'
    assert chrome.label(None, 3, 9, template=True) == '03'
    s = {'number': 'n_total', 'main_slides': 15}
    assert chrome.label(s, 14, 20) == '14 / 15' and chrome.label(s, 16, 20) == '16'
    assert chrome.label(s, 2, 20, template=True) == '02 / 15'
    assert chrome.label({'number': 'n'}, 7, 9) == '7'
    assert chrome.label({'number': 'none'}, 7, 9) is None
    assert chrome.shows_on(None, 'title') is False and chrome.shows_on({'on_cover': True}, 'title') is True


def test_progress_geometry_bar_segments_and_appendix():
    bar = chrome.progress_frames({'progress': 'bar', 'main_slides': 4}, 1, 6, 1280, 720)
    assert bar == [('track', 0.0, 714.0, 1280, 6.0), ('fill', 0.0, 714.0, 320.0, 6.0)]
    top = chrome.progress_frames({'progress': 'bar', 'position': 'top', 'thickness': 10}, 2, 2, 1920, 1080)
    assert top[0][2] == 0.0 and top[0][4] == 15.0 and top[1][3] == 1920
    seg = chrome.progress_frames({'progress': 'segments'}, 2, 5, 1280, 720)
    assert [r for r, *_ in seg] == ['fill', 'fill', 'track', 'track', 'track']
    assert seg[-1][1] + seg[-1][3] == pytest.approx(1280)
    assert chrome.progress_frames({'progress': 'bar', 'main_slides': 4}, 5, 6, 1280, 720) == []
    assert chrome.progress_frames(None, 1, 6, 1280, 720) == []


def _plan(slide, settings, number=3, total=5):
    return build_plan(slide, deck_title='Deck', theme=theme_of({'theme': 'claro'}), number=number, total=total, chrome=settings)


def test_plan_without_settings_is_unchanged_and_settings_drive_footer_and_progress():
    slide = {'id': 's', 'layout': 'bullets', 'title': 'T', 'blocks': [{'type': 'bullets', 'items': ['a']}]}
    plain = _plan(slide, None)
    assert [b.paras[0].text for b in plain.boxes if b.role == 'footer'] == ['Deck', '3 / 5']
    assert not [b for b in plain.boxes if b.role == 'progress']
    custom = _plan(slide, {'number': 'n', 'progress': 'bar', 'show_title': False})
    assert [b.paras[0].text for b in custom.boxes if b.role == 'footer'] == ['3']
    fill = [b for b in custom.boxes if b.role == 'progress' and b.fill == 'accent']
    assert len(fill) == 1 and fill[0].w == pytest.approx(1280 * 3 / 5)
    cover = {'id': 'c', 'layout': 'title', 'title': 'Cover', 'blocks': []}
    assert not [b for b in _plan(cover, {'progress': 'bar'}).boxes if b.role in ('progress', 'footer')]
    shown = _plan(cover, {'progress': 'bar', 'on_cover': True})
    assert [b.paras[0].text for b in shown.boxes if b.role == 'footer'] == ['3 / 5']
    assert [b for b in shown.boxes if b.role == 'progress']


def test_html_and_pptx_draw_the_same_settings():
    deck = sample_deck()
    deck['chrome'] = {'number': 'n_total', 'progress': 'segments', 'main_slides': len(deck['slides']) - 1}
    html, _ = render_document(deck, assets=NO_ASSETS, inline=True, mode='stacked')
    assert html.count('background:var(--accent)') > len(deck['slides'])
    prs = Presentation(io.BytesIO(build_pptx(deck, assets=NO_ASSETS)))
    texts = [sh.text for sh in prs.slides[2].shapes if sh.has_text_frame]
    assert f'3 / {len(deck["slides"]) - 1}' in texts
    last = [sh.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
    assert str(len(deck['slides'])) in last or deck['slides'][-1]['layout'] in ('closing', 'title', 'section')


def test_tool_roundtrip_validation_and_reset(services):
    d = decks.create_deck(services, DeckCreate(title='Defensa', slide_count=3))
    for t in ('Uno', 'Dos', 'Tres'):
        decks.add_slide(services, d['id'], title=t, notes='n')
    got = agent_tools.call_tool(services, 'deck_chrome', {'deck_id': d['id']})
    assert got['custom'] is False and got['settings']['number'] == 'auto'
    assert 'chrome' not in decks.deck_view(services, d['id'])
    with pytest.raises(CiceroError):
        agent_tools.call_tool(services, 'deck_chrome', {'deck_id': d['id'], 'action': 'set', 'settings': {'main_slides': 9}})
    done = agent_tools.call_tool(services, 'deck_chrome', {'deck_id': d['id'], 'action': 'set',
                                                            'settings': {'number': 'n_total', 'progress': 'bar', 'main_slides': 2}})
    assert done['custom'] and decks.deck_view(services, d['id'])['chrome']['main_slides'] == 2
    exported = agent_tools.call_tool(services, 'deck_export', {'deck_id': d['id'], 'format': 'pptx'})
    prs = Presentation(exported['path'])
    assert '1 / 2' in [s.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert '3' in [s.text for s in prs.slides[2].shapes if s.has_text_frame]
    assert agent_tools.call_tool(services, 'deck_chrome', {'deck_id': d['id'], 'action': 'reset'})['custom'] is False


def test_template_number_and_progress(services, tmp_path):
    from pptx.util import Pt
    file = tmp_path / 'numbered.pptx'
    template(file)
    prs = Presentation(file)
    box = prs.slides[0].shapes.add_textbox(Inches(14), Inches(8.3), Inches(1.5), Inches(.5))
    box.name = 'Slide Number Placeholder 9'; box.text = '00'; box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(file)
    d = decks.create_deck(services, DeckCreate(title='Plantilla', slide_count=3))
    for t in ('Uno', 'Dos', 'Tres'):
        decks.add_slide(services, d['id'], title=t, notes='n')
    agent_tools.call_tool(services, 'deck_template', {'deck_id': d['id'], 'path': str(file)})
    before = Presentation(agent_tools.call_tool(services, 'deck_export', {'deck_id': d['id'], 'format': 'pptx'})['path'])
    assert '02' in [s.text for s in before.slides[1].shapes if s.has_text_frame]
    agent_tools.call_tool(services, 'deck_chrome', {'deck_id': d['id'], 'action': 'set',
                                                     'settings': {'number': 'n_total', 'progress': 'bar', 'main_slides': 2}})
    after = Presentation(agent_tools.call_tool(services, 'deck_export', {'deck_id': d['id'], 'format': 'pptx'})['path'])
    assert '02 / 02' in [s.text for s in after.slides[1].shapes if s.has_text_frame]
    assert '03' in [s.text for s in after.slides[2].shapes if s.has_text_frame]
    bars = [s for s in after.slides[0].shapes if s.name == 'Decoration' and s.top > after.slide_height * 0.95]
    assert len(bars) == 2 and max(b.width for b in bars) == after.slide_width
    assert not [s for s in after.slides[2].shapes if s.name == 'Decoration' and s.top > after.slide_height * 0.95]
