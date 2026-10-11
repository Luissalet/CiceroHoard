import io
from concurrent.futures import ThreadPoolExecutor
import pytest
from openpyxl import Workbook
from cicero_hoard import datasets as ds,decks
from cicero_hoard.agent_tools import call_tool,TOOLS_BY_NAME
from cicero_hoard.errors import CiceroError,Refused,RevisionConflict
from cicero_hoard.hoard_link.agentkit import AppError
from conftest import new_deck

@pytest.mark.parametrize('raw,mode,expected',[
 ('1.234,56','es',1234.56),('1,234.56','en',1234.56),('1.234,56','auto',1234.56),
 ('3,49','auto',3.49),('-3.49','auto',-3.49),('1.234','auto',None),('1,234','auto',None),
 ('12,34,56','en',None),('12.34.56','es',None),(True,'auto',None),(float('nan'),'auto',None),
 ('1e3','auto',1000),('0','auto',0)])
def test_numeric_profiles_do_not_guess_ambiguous_values(raw,mode,expected):
 assert ds.numeric(raw,mode)==expected

def imported(services,**kw):
 deck=new_deck(services)
 args={'deck_id':deck['id'],'text':'Region;Value;Other\nNorth;1.234,56;3,49\nNorth;2,5;\nSouth;-4,2;6,03','number_format':'es',**kw}
 return deck,call_tool(services,'dataset_import',args)

def test_snapshot_mapping_citations_retries_and_undo(services):
 deck,data=imported(services)
 again=call_tool(services,'dataset_import',{'deck_id':deck['id'],'text':'Region;Value;Other\nNorth;1.234,56;3,49\nNorth;2,5;\nSouth;-4,2;6,03','number_format':'es'})
 assert again['reused'] and again['source_id']==data['source_id']
 assert ds.undo(services,{'etag':'reused'})=={'unchanged':True}
 result=call_tool(services,'dataset_visual',{'deck_id':deck['id'],'source_id':data['id'],'category_column':'Region','series_columns':['Value','Other']})
 assert result['block']['series'][0]['values']==[1234.56,2.5,-4.2]
 assert result['block']['series'][1]['values']==[3.49,None,6.03]
 slide=decks.add_slide(services,deck['id'],blocks=[{'type':'text','text':'Keep'}],notes='Private note')
 args={'deck_id':deck['id'],'slide_id':slide['id'],'block_index':1,'block':result['block'],'source_ids':result['source_ids'],'expected_revision':1}
 updated=call_tool(services,'slide_visual_update',args)
 assert updated['notes']=='Private note' and updated['sources']==[data['id']] and updated['blocks'][0]['text']=='Keep'
 with pytest.raises(RevisionConflict):call_tool(services,'slide_visual_update',args)
 args['expected_revision']=2
 assert call_tool(services,'slide_visual_update',args)['changed'] is False
 record=ds.track(services,ds.ImportArgs(deck_id=deck['id']),data)
 with pytest.raises(AppError):ds.undo(services,record)

def test_transform_preserves_parent_and_records_definition(services):
 deck,data=imported(services)
 derived=call_tool(services,'dataset_transform',{'deck_id':deck['id'],'source_id':data['id'],'group_by':['Region'],
   'aggregations':[{'column':'Value','operation':'mean','name':'Average'},{'column':'Region','operation':'count','name':'Rows'}]})
 assert derived['rows']==[['North',pytest.approx(618.53),2],['South',-4.2,1]]
 assert derived['origin']['derived_from']==data['id']
 assert ds.load(services,deck['id'],data['id'])['rows'][0][1]=='1.234,56'
 mapped=call_tool(services,'dataset_visual',{'deck_id':deck['id'],'source_id':derived['id'],'category_column':'Region','series_columns':['Average']})
 assert mapped['source_ids']==[derived['id']]

def test_sort_handles_mixed_types_and_keeps_missing_last(services):
 deck,data=imported(services,text='Key,Value\na,2\nb,text\nc,\nd,1',number_format='auto')
 out=ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],sort_column='Value',descending=True))
 assert out['rows'][-1]==['c',''] and len(out['rows'])==4
 out=ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],aggregations=[{'column':'Key','operation':'last'}]))
 assert out['rows']==[['d']]

def test_xlsx_sheet_dates_uncached_formulas_and_upload(client):
 deck=new_deck(client.svc);book=Workbook();book.active.title='Ignore';sheet=book.create_sheet('Measurements')
 sheet.append(['Name','Value']);sheet.append(['Observed',2.01]);sheet.append(['Formula','=SUM(1,2)']);buffer=io.BytesIO();book.save(buffer)
 response=client.post(f"/api/decks/{deck['id']}/datasets",files={'file':('data.xlsx',buffer.getvalue())},data={'sheet':'Measurements'})
 assert response.status_code==201
 data=response.json();assert data['rows']==[['Observed',2.01],['Formula',None]]
 assert data['origin']['sheets']==['Ignore','Measurements'] and 'Uncached formula: B3' in data['warnings']
 assert client.get(f"/api/decks/{deck['id']}/datasets/{data['id']}?limit=1").json()['has_more']

def test_import_retry_concurrency_has_one_source(services):
 deck=new_deck(services);args=ds.ImportArgs(deck_id=deck['id'],text='a,b\nx,2')
 with ThreadPoolExecutor(max_workers=4) as pool:results=list(pool.map(lambda _:ds.import_data(services,args),range(8)))
 assert len({r['id'] for r in results})==1 and sum(not r['reused'] for r in results)==1

def test_ambiguous_cells_must_be_resolved_explicitly(services):
 deck,data=imported(services,text='Name;Value\nA;1.234',number_format='auto')
 assert data['warnings']
 with pytest.raises(CiceroError):ds.visual(services,ds.VisualArgs(deck_id=deck['id'],source_id=data['id'],category_column='Name',series_columns=['Value']))

def test_dataset_tool_has_accountable_draft_writes(services):
 deck,data=imported(services)
 for name in ('dataset_import','dataset_transform'):
  assert TOOLS_BY_NAME[name].track and TOOLS_BY_NAME[name].undo and TOOLS_BY_NAME[name].annotations['draftSafeHint']
 track=ds.track(services,ds.ImportArgs(deck_id=deck['id']),data)
 assert ds.undo(services,track,dry_run=True)=={'would_remove':data['id']}
 assert ds.undo(services,track)=={'removed':data['id']}

def test_invalid_grouping_output_and_numeric_comparisons(services):
 deck,data=imported(services,text='Key,Value\na,text',number_format='auto')
 with pytest.raises(CiceroError):ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],group_by=['Key']))
 with pytest.raises(CiceroError):ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],filters=[{'column':'Value','operator':'gt','value':0}]))

def test_invalid_workbook_is_a_clear_import_error(client):
 deck=new_deck(client.svc)
 response=client.post(f"/api/decks/{deck['id']}/datasets",files={'file':('bad.xlsx',b'not a workbook')})
 assert response.status_code==400 and 'workbook' in response.json()['error']

def test_agent_mapping_keeps_complete_large_numeric_arrays(services):
 deck,data=imported(services,text='Name,Value\n'+'\n'.join(f'Category_{i:04},'+str(i) for i in range(2000)),number_format='auto')
 result=call_tool(services,'dataset_visual',{'deck_id':deck['id'],'source_id':data['id'],'category_column':'Name','series_columns':['Value'],'limit':2000},cap=True)
 assert len(result['block']['categories'])==2000 and result['block']['series'][0]['values']==list(range(2000))
 assert not result.get('truncated')

@pytest.mark.parametrize('raw,mode,expected',[(',5','es',.5),('-,5','es',-.5),('+,5','auto',.5),('.5','en',.5)])
def test_leading_decimal_separator_is_valid(raw,mode,expected):
 assert ds.numeric(raw,mode)==expected

def test_table_maps_ten_columns_but_chart_still_checks_six(client):
 deck=new_deck(client.svc)
 data=client.post(f"/api/decks/{deck['id']}/datasets",json={'text':','.join('C'+str(i) for i in range(10))+'\n'+','.join(str(i) for i in range(10))}).json()
 base=f"/api/decks/{deck['id']}/datasets/{data['id']}/visual"
 result=client.post(base,json={'kind':'table','series_columns':data['columns']})
 assert result.status_code==200 and len(result.json()['block']['headers'])==10
 assert client.post(base,json={'kind':'chart','category_column':'C0','series_columns':data['columns']}).status_code==400

def test_undo_preserves_cited_derived_ancestors_and_handles_reverse_order(services):
 deck,parent=imported(services)
 args=ds.TransformArgs(deck_id=deck['id'],source_id=parent['id'],aggregations=[{'column':'Value','operation':'sum'}])
 child=ds.transform(services,args)
 slide=decks.add_slide(services,deck['id'],sources=[child['id']],notes='Derived evidence')
 record=ds.track(services,ds.ImportArgs(deck_id=deck['id']),parent)
 with pytest.raises(AppError,match='Derived datasets'):ds.undo(services,record)
 assert ds.load(services,deck['id'],parent['id'])['rows'] and ds.load(services,deck['id'],child['id'])['origin']['derived_from']==parent['id']
 decks.delete_slide(services,deck['id'],slide['id'])
 ds.undo(services,ds.track(services,args,child))
 assert ds.undo(services,record)=={'removed':parent['id']}

def test_transform_cannot_insert_after_parent_removed(services,monkeypatch):
 deck,parent=imported(services);snapshot=ds._snapshot
 def interleave(svc,deck_id,title,payload):
  decks.remove_source(svc,deck_id,parent['id'])
  return snapshot(svc,deck_id,title,payload)
 monkeypatch.setattr(ds,'_snapshot',interleave)
 with pytest.raises(CiceroError,match='parent dataset was removed'):
  ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=parent['id'],sort_column='Value'))
 assert not decks.deck_view(services,deck['id'])['sources']

def test_contains_matches_zero_without_treating_it_as_empty(services):
 deck,data=imported(services,text='Name,Value\nA,0\nB,1',number_format='auto')
 result=ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],filters=[{'column':'Value','operator':'contains','value':'0'}]))
 assert result['rows']==[['A','0']]

def test_dataset_chain_whole_session_undo_preview_and_execution(client):
 from test_agent_accountability import agent_call,ok,undo
 deck=new_deck(client.svc);session='data-chain'
 parent=ok(agent_call(client,'dataset_import',{'deck_id':deck['id'],'text':'Name,Value\nA,2\nB,3'},session=session))
 child=ok(agent_call(client,'dataset_transform',{'deck_id':deck['id'],'source_id':parent['id'],
     'aggregations':[{'column':'Value','operation':'sum'}]},session=session))
 ok(agent_call(client,'slide_add',{'deck_id':deck['id'],'title':'Evidence','sources':[child['id']]},session=session))
 preview=ok(undo(client,session,dry_run=True))
 assert preview['complete'] and len(preview['would_undo'])==3
 result=ok(undo(client,session,confirm=True))
 assert result['complete'] and not decks.deck_view(client.svc,deck['id'])['sources'] and not decks.deck_view(client.svc,deck['id'])['slides']

def test_overflowing_aggregate_never_saves_nonfinite_snapshot(services):
 deck,data=imported(services,text='Name,Value\nA,1e308\nB,1e308',number_format='auto')
 with pytest.raises(ValueError):ds.transform(services,ds.TransformArgs(deck_id=deck['id'],source_id=data['id'],aggregations=[{'column':'Value','operation':'sum'}]))
 assert len(decks.deck_view(services,deck['id'])['sources'])==1

def test_session_undo_preview_preserves_external_parent_citation(client):
 from test_agent_accountability import agent_call,ok,undo
 deck=new_deck(client.svc);session='data-chain-external'
 parent=ok(agent_call(client,'dataset_import',{'deck_id':deck['id'],'text':'Name,Value\nA,2'},session=session))
 child=ok(agent_call(client,'dataset_transform',{'deck_id':deck['id'],'source_id':parent['id'],'sort_column':'Value'},session=session))
 ok(agent_call(client,'slide_add',{'deck_id':deck['id'],'title':'Derived','sources':[child['id']]},session=session))
 external=decks.add_slide(client.svc,deck['id'],title='External evidence',sources=[parent['id']])
 preview=ok(undo(client,session,dry_run=True))
 assert not preview['complete'] and len(preview['conflicts'])==1
 result=ok(undo(client,session,confirm=True))
 assert not result['complete'] and len(result['conflicts'])==1
 current=decks.deck_view(client.svc,deck['id'])
 assert [s['id'] for s in current['sources']]==[parent['id']]
 assert [s['id'] for s in current['slides']]==[external['id']]
