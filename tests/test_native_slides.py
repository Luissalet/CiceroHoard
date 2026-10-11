import io
import json
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from cicero_hoard import decks, native_slides
from cicero_hoard.agent_tools import call_tool
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.layouts import count_lines
from conftest import new_deck, png_bytes
from sample_deck import sample_deck
from test_templates import template


def test_chart_ids_work_with_native_signed_integer_readers():
    with zipfile.ZipFile(io.BytesIO(build_pptx(sample_deck(), assets=lambda _:None))) as package:
        for name in package.namelist():
            if name.startswith('ppt/charts/chart') and name.endswith('.xml'):
                root=etree.fromstring(package.read(name))
                ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                ids=[int(e.get('val')) for e in root.findall('.//c:axId',ns)]
                refs=[int(e.get('val')) for e in root.findall('.//c:crossAx',ns)]
                assert all(0<value<2**31 for value in ids)
                assert set(refs)==set(ids)


def test_plain_text_cancels_imported_hanging_indents():
    prs=Presentation(io.BytesIO(build_pptx(sample_deck(),assets=lambda _:None)))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    ppr=paragraph._p.pPr
                    if ppr is not None and ppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buNone') is not None:
                        assert ppr.get('marL')=='0' and ppr.get('indent')=='0'


def test_explicit_breaks_contribute_to_fit():
    assert count_lines('one\ntwo\nthree',1000,32,'sans',False)==3
    assert count_lines('one\vthree',1000,32,'sans',False)==2
    assert count_lines('one\n\nthree',1000,32,'sans',False)==3


def test_template_export_drops_only_donor_library_metadata():
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from cicero_hoard.template_pptx import _drop_library_metadata
    prs=Presentation()
    metadata=Part(PackURI('/customXml/item1.xml'),'application/xml',prs.part.package,
                  b'<properties xmlns="urn:example/office/2006/metadata/properties"/>')
    custom=Part(PackURI('/customXml/item2.xml'),'application/xml',prs.part.package,
                b'<properties xmlns="urn:custom:data">Keep this data</properties>')
    removed=prs.part.relate_to(metadata,RT.CUSTOM_XML)
    kept=prs.part.relate_to(custom,RT.CUSTOM_XML)
    _drop_library_metadata(prs)
    assert removed not in prs.part.rels and kept in prs.part.rels


def test_template_frame_overrides_are_preserved_and_validated(services,tmp_path):
    file=tmp_path/'brand.pptx';binding=template(file)
    deck=new_deck(services)
    decks.add_slide(services,deck['id'],title='Expanded title',layout='chart',notes='Notes',
                    blocks=[{'type':'chart','chart':'bar','categories':['A','B'],'series':[{'name':'Test','values':[3.49,4.26]}]}])
    binding.update(title_frame=[96,48,1200,100],title_font_points=44,chart_font_points=18)
    call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file),'bindings':{'chart':binding}})
    prs=Presentation(call_tool(services,'deck_export',{'deck_id':deck['id'],'format':'pptx'})['path'])
    title=next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text=='Expanded title')
    assert title.width==1200*9525 and title.top==48*9525
    assert title.text_frame.paragraphs[0].runs[0].font.size==Pt(44)
    chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert chart.category_axis.tick_labels.font.size==Pt(18)
    binding['title_frame']=[0,0,float('nan'),100]
    with pytest.raises(Exception):
        call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file),'bindings':{'chart':binding}})


def test_native_preview_and_pdf_use_the_same_render(services,tmp_path):
    file=tmp_path/'brand.pptx';template(file)
    deck=new_deck(services);slide=decks.add_slide(services,deck['id'],title='Brand',notes='Private notes')
    call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file)})
    calls=[]
    def fake(data,folder):
        calls.append(data);folder.mkdir(parents=True,exist_ok=True)
        (folder/'slides.pdf').write_bytes(b'%PDF-native')
        (folder/'slide-1.png').write_bytes(png_bytes())
        report={'renderer':'native_office','slide_count':1,'slides':[{'position':1,'shapes':[]}]}
        (folder/'report.json').write_text(json.dumps(report))
        return report
    services._native_render_fn=fake
    preview=services.preview_html(deck['id'],slide['id'])
    assert 'data:image/png;base64,' in preview and 'Private notes' not in preview
    out=services.export(deck['id'],'pdf')
    assert decks.export_file(services,out['id'])[0].read_bytes()==b'%PDF-native'
    assert services.check(deck['id'])['native_review']['status']=='rendered'
    page=services.export(deck['id'],'html')
    assert 'new Reveal' in decks.export_file(services,page['id'])[0].read_text()
    assert all(data.startswith(b'PK') for data in calls)


def test_native_text_bounds_detect_actual_overflow():
    report={'slides':[{'position':2,'shapes':[{'name':'Body','text':'Measured',
        'frame':[10,20,100,30],'text_bounds':[10,20,100,80]}]}]}
    result=native_slides.layout_findings(report)
    assert len(result)==1 and result[0]['position']==2 and result[0]['kind']=='native_text_overflow'
    report['slides'][0]['shapes'][0]['text_bounds']=[10,20,100,30]
    assert native_slides.layout_findings(report)==[]


def test_native_cache_returns_without_launching_an_engine(tmp_path,monkeypatch):
    monkeypatch.setattr(native_slides,'available',lambda:True)
    (tmp_path/'slides.pdf').write_bytes(b'%PDF')
    (tmp_path/'report.json').write_text(json.dumps({'slide_count':1}))
    monkeypatch.setattr(native_slides.proc,'run',lambda *a,**k:pytest.fail('must use cached rendering'))
    assert native_slides.render(b'PK',tmp_path)=={'slide_count':1}

def test_live_template_resolves_scheme_colors_and_preserves_line_breaks(services,tmp_path):
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.oxml.ns import qn
    from cicero_hoard.template_live import scenes
    file=tmp_path/'brand.pptx';binding=template(file)
    prs=Presentation(str(file));source=prs.slides[binding['source_slide']-1]
    shape=next(s for s in source.shapes if s.shape_id==binding['title_shape'])
    shape.text_frame.paragraphs[0].runs[0].font.color.theme_color=MSO_THEME_COLOR.ACCENT_4
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part=next(r.target_part for r in source.slide_layout.slide_master.part.rels.values() if r.reltype==RT.THEME)
    root=etree.fromstring(part.blob);slot=root.find('.//'+qn('a:accent4'))
    slot[0].set('val','FFFFFF');part._blob=etree.tostring(root);prs.save(str(file))
    deck=new_deck(services);decks.add_slide(services,deck['id'],title='Line one\nLine two',layout='chart')
    call_tool(services,'deck_template',{'deck_id':deck['id'],'path':str(file),'bindings':{'chart':binding}})
    data=decks.deck_view(services,deck['id']);scene,_,_=scenes(data,services.assets_lookup())
    assert scene[0]['theme']['colors']['title_ink']=='#FFFFFF'
    assert any(p.text=='Line one\nLine two' for b in scene[0]['boxes'] for p in b.paras)
