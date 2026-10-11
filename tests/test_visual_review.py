"""Visual review: rule bookkeeping with a fake measurement, the allowed fix, the tool and routes, and one real browser run."""

from __future__ import annotations

import pytest

from cicero_hoard import decks, visual_review as vr
from cicero_hoard.agent_tools import call_tool
from cicero_hoard.errors import PdfUnavailable
from conftest import new_deck


def fake_rows(found_by_index):
    def measure(deck, *, assets, png_dir=None):
        rows = []
        for i, s in enumerate(deck["slides"]):
            rows.append({"id": s["id"], "found": [dict(f) for f in found_by_index.get(i, [])], "hasTitle": bool(s.get("title")),
                         "hasChart": False, "hasImage": False})
        return rows
    return measure


def test_review_classifies_and_adds_layout_checks():
    deck = {"slides": [{"id": "a", "title": "T", "layout": "chart", "blocks": []},
                       {"id": "b", "title": "", "layout": "bullets", "blocks": []}]}
    out = vr.review(deck, assets=lambda _: None, measure_fn=fake_rows({1: [{"rule": "contrast", "where": "x", "detail": "2.1:1 < 4.5:1"},
                                                                           {"rule": "overlap", "where": "a / b", "detail": "30%"}]}))
    assert not out["ok"] and out["errors"] == 1 and out["warnings"] == 2
    assert [f["rule"] for f in out["slides"][0]["findings"]] == ["missing_chart"]
    contrast = out["slides"][1]["findings"][0]
    assert contrast["severity"] == "error" and contrast["fix"] and contrast["auto_fix"] is False


def test_shrink_patch_only_touches_text_without_a_smaller_size():
    slide = {"blocks": [{"type": "text", "text": "x"}, {"type": "bullets", "items": ["a"], "font_size": 14},
                        {"type": "image", "asset_id": "z"}]}
    blocks = vr.shrink_patch(slide)
    assert blocks[0]["font_size"] == 14 and blocks[1]["font_size"] == 14 and "font_size" not in blocks[2]
    assert vr.shrink_patch({"blocks": [{"type": "text", "text": "x", "font_size": 14}]}) is None


def test_fix_applies_minimum_font_and_reviews_again(services):
    deck = new_deck(services)
    slide = decks.add_slide(services, deck["id"], title="Long", blocks=[{"type": "text", "text": "Mucho texto"}])
    calls = []

    def measure(d, *, assets, png_dir=None):
        calls.append(1)
        s = d["slides"][0]
        small = (s["blocks"][0].get("font_size") == vr.MIN_FONT)
        return [{"id": s["id"], "found": [] if small else [{"rule": "text_overflow", "where": "Mucho", "detail": "needs 900px", "need": 900, "have": 500}],
                 "hasTitle": True, "hasChart": False, "hasImage": False}]

    services._visual_measure_fn = measure
    first = call_tool(services, "deck_visual_review", {"deck_id": deck["id"]})
    assert not first["ok"] and first["slides"][0]["findings"][0]["auto_fix"] is True
    fixed = call_tool(services, "deck_visual_review", {"deck_id": deck["id"], "fix": True})
    assert fixed["ok"] and fixed["fixed"][0]["slide_id"] == slide["id"] and fixed["before"]["errors"] == 1
    assert decks.get_slide(services, deck["id"], slide["id"])["blocks"][0]["font_size"] == vr.MIN_FONT
    assert len(calls) == 3


def test_real_browser_review_writes_pngs_and_finds_real_problems(services, client):
    deck = new_deck(services)
    decks.add_slide(services, deck["id"], title="Fine", blocks=[{"type": "bullets", "items": ["Uno", "Dos"]}])
    decks.add_slide(services, deck["id"], title="Crowded", layout="freeform",
                    blocks=[{"type": "text", "text": "palabra " * 300, "frame": [0, 0, 0.2, 0.08]},
                            {"type": "text", "text": "encima", "frame": [0.1, 0.0, 0.5, 0.3]},
                            {"type": "text", "text": "fuera", "frame": [0.9, 0.5, 0.3, 0.1]}])
    try:
        out = services.visual_review(deck["id"])
    except PdfUnavailable:
        pytest.skip("no browser here")
    assert len(out["slides"]) == 2 and all(s["png_url"].endswith(".png") for s in out["slides"])
    assert not [f for f in out["slides"][0]["findings"] if f["severity"] == "error"]
    rules = {f["rule"] for f in out["slides"][1]["findings"]}
    # the long text grew past the bottom: overflow with the automatic fix; the box past the right edge is off the slide
    assert {"off_slide", "overlap", "text_overflow"} <= rules, out["slides"][1]["findings"]
    grown = next(f for f in out["slides"][1]["findings"] if f["rule"] == "text_overflow")
    assert grown["auto_fix"] and "the slide leaves" in grown["detail"]
    off = [f for f in out["slides"][1]["findings"] if f["rule"] == "off_slide"]
    assert [f["where"] for f in off] == ["fuera"], off
    assert (services.visual_review_dir(deck["id"]) / "slide-01.png").is_file()


def test_routes(client):
    d = client.post("/api/decks", json={"title": "HTTP review"}).json()
    client.post(f"/api/decks/{d['id']}/slides", json={"title": "A", "blocks": [{"type": "text", "text": "x"}]})
    try:
        out = client.post(f"/api/decks/{d['id']}/visual-review", json={}).json()
    except PdfUnavailable:
        pytest.skip("no browser here")
    if "slides" not in out:
        pytest.skip(out.get("error", "no browser here"))
    png = client.get(out["slides"][0]["png_url"])
    assert png.status_code == 200 and png.headers["content-type"] == "image/png" and png.content[:4] == b"\x89PNG"
    assert client.get(f"/api/decks/{d['id']}/visual-review/../../secret").status_code == 404
    assert client.get(f"/api/decks/{d['id']}/visual-review/slide-99.png").status_code == 404


def test_fix_never_overwrites_an_edit_made_while_measuring(services):
    """The automatic fix applies only over the revision that was measured: a slide edited during the measurement keeps
    the edit and is reported as skipped; the other slides are still fixed, and the deck is reviewed again."""
    deck = new_deck(services)
    a = decks.add_slide(services, deck["id"], title="Edited meanwhile", blocks=[{"type": "text", "text": "Viejo"}])
    b = decks.add_slide(services, deck["id"], title="Untouched", blocks=[{"type": "text", "text": "Mucho texto"}])
    state = {"calls": 0}

    def measure(d, *, assets, png_dir=None, **_):
        state["calls"] += 1
        if state["calls"] == 1:  # the user edits slide A while the browser is measuring
            decks.patch_slide(services, deck["id"], a["id"], {"blocks": [{"type": "text", "text": "Nuevo", "font_size": 30}]})
        rows = []
        for s in d["slides"]:
            small = s["blocks"][0].get("font_size") in (vr.MIN_FONT, 30)
            rows.append({"id": s["id"], "found": [] if small else [{"rule": "text_overflow", "where": "x", "detail": "needs 900px",
                                                                   "need": 900, "have": 500}],
                         "hasTitle": True, "hasChart": False, "hasImage": False})
        return rows

    services._visual_measure_fn = measure
    out = services.visual_review(deck["id"], fix=True)
    assert [f["slide_id"] for f in out["fixed"]] == [b["id"]]
    assert [s["slide_id"] for s in out["skipped"]] == [a["id"]]
    kept = decks.get_slide(services, deck["id"], a["id"])["blocks"][0]
    assert kept["text"] == "Nuevo" and kept["font_size"] == 30  # the user's edit survives
    assert decks.get_slide(services, deck["id"], b["id"])["blocks"][0]["font_size"] == vr.MIN_FONT
    assert out["ok"] and state["calls"] == 2


def test_template_deck_is_measured_in_the_scene_the_preview_shows(services, tmp_path):
    """A template deck is reviewed in its real scene: its own slide size and the template's placement, so a title the
    template puts past the right edge is reported, and the capture is the preview's picture."""
    import io

    from pptx import Presentation
    from pptx.util import Inches, Pt

    from cicero_hoard import agent_tools
    from conftest import png_bytes

    path = tmp_path / "wide-title.pptx"
    p = Presentation()
    p.slide_width, p.slide_height = Inches(16), Inches(9)
    s = p.slides.add_slide(p.slide_layouts[6])
    title = s.shapes.add_textbox(Inches(12), Inches(1), Inches(8), Inches(1.5))
    title.text = "Example title"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    p.save(path)

    def native_render(data, folder):  # the template furniture as pictures, without an office suite
        folder.mkdir(parents=True, exist_ok=True)
        for n in range(1, len(Presentation(io.BytesIO(data)).slides) + 1):
            (folder / f"slide-{n}.png").write_bytes(png_bytes())
        return {"ok": True}

    services._native_render_fn = native_render
    deck = new_deck(services)
    decks.add_slide(services, deck["id"], title="A title the template places past the edge",
                    blocks=[{"type": "bullets", "items": ["Uno", "Dos"]}])
    agent_tools.call_tool(services, "deck_template", {"deck_id": deck["id"], "path": str(path), "bindings": {"bullets": {
        "source_slide": 1, "title_shape": title.shape_id, "keep_text_shapes": [], "content_frame": [96, 300, 960, 300]}}})
    try:
        out = services.visual_review(deck["id"])
    except PdfUnavailable:
        pytest.skip("no browser here")
    off = [f for f in out["slides"][0]["findings"] if f["rule"] == "off_slide"]
    assert off and "A title the template places" in off[0]["where"], out["slides"][0]["findings"]
    assert (services.visual_review_dir(deck["id"]) / "slide-01.png").is_file()
