import io
import pytest
from pptx import Presentation
from pptx.enum.action import PP_ACTION
from pydantic import ValidationError
from cicero_hoard import decks,navigation,presentation
from cicero_hoard.agent_tools import call_tool
from cicero_hoard.models import NavigationAction
from cicero_hoard.errors import CiceroError,RevisionConflict
from cicero_hoard.export_pptx import build_pptx
from cicero_hoard.render import render_document,present_document
from conftest import new_deck

def route_deck(svc):
    d=new_deck(svc)
    first=decks.add_slide(svc,d['id'],title='Menu',blocks=[{'type':'text','text':'Details','build':1,'id':'details'}],notes='PRIVATE')
    detail=decks.add_slide(svc,d['id'],title='Detail',blocks=[{'type':'text','text':'Return','action':{'kind':'return'}}])
    return d,first,detail

def test_action_schema_validates_destinations_and_active_url_schemes():
    for action in [{'kind':'slide'},{'kind':'url','url':'javascript:alert(1)'},{'kind':'url','url':'https://'},
                   {'kind':'return','slide_id':'x'},{'kind':'slide','slide_id':'s','url':'https://example.org'}]:
        with pytest.raises(ValidationError):NavigationAction.model_validate(action)
    assert NavigationAction(kind='url',url='mailto:test@example.org').url

def test_mcp_link_preserves_notes_revisions_identity_and_reports_deleted_target(services):
    d,first,detail=route_deck(services)
    args={'deck_id':d['id'],'slide_id':first['id'],'block_index':0,'action':'set',
          'target':{'kind':'slide','slide_id':detail['id'],'label':'Details'},'expected_revision':first['revision']}
    result=call_tool(services,'deck_navigation',args)
    assert result['notes']=='PRIVATE' and result['blocks'][0]['id']=='details'
    with pytest.raises(RevisionConflict):call_tool(services,'deck_navigation',args)
    decks.reorder_slides(services,d['id'],[detail['id'],first['id']])
    out=call_tool(services,'deck_navigation',{'deck_id':d['id']})
    assert out['links'][1]['target']['slide_id']==detail['id'] and out['links'][1]['valid']
    decks.delete_slide(services,d['id'],detail['id'])
    assert not call_tool(services,'deck_navigation',{'deck_id':d['id']})['links'][0]['valid']
    assert 'link_missing' in {i['kind'] for i in services.check(d['id'])['issues']}
    call_tool(services,'deck_navigation',{**args,'action':'clear','expected_revision':result['revision']})
    assert not call_tool(services,'deck_navigation',{'deck_id':d['id']})['links']

def test_live_nested_branch_return_restores_click_stage_and_does_not_mutate_notes(services):
    d,first,detail=route_deck(services)
    presentation.start(services,presentation.StartArgs(deck_id=d['id']))
    control=lambda **kw:presentation.control(services,presentation.ControlArgs(deck_id=d['id'],**kw))
    control(action='next')
    branched=control(action='branch',slide_id=detail['id'])
    assert branched['slide_index']==1 and branched['return_depth']==1
    with pytest.raises(RevisionConflict):control(action='return',expected_version=1)
    control(action='branch',slide_id=first['id'])
    assert control(action='return')['slide_index']==1
    back=control(action='return')
    assert back['slide_index']==0 and back['fragment_index']==0 and back['return_depth']==0
    with pytest.raises(CiceroError):control(action='return')
    assert decks.get_slide(services,d['id'],first['id'])['notes']=='PRIVATE'

def test_native_internal_url_and_return_links_survive_reorder_and_motion(services):
    d,first,detail=route_deck(services)
    navigation.run(services,navigation.NavigationArgs(deck_id=d['id'],action='set',slide_id=first['id'],block_index=0,target=NavigationAction(kind='slide',slide_id=detail['id'])))
    decks.add_slide(services,d['id'],title='URL',blocks=[{'type':'text','text':'External','action':{'kind':'url','url':'https://example.org/a?q=1&v=2'}},
                       {'type':'text','text':'Next','action':{'kind':'next'}},{'type':'text','text':'Previous','action':{'kind':'previous'}}])
    deck=decks.deck_view(services,d['id'])
    prs=Presentation(io.BytesIO(build_pptx(deck,assets=services.assets_lookup())))
    linked=[s for s in prs.slides[0].shapes if s.click_action.action==PP_ACTION.NAMED_SLIDE]
    assert linked and linked[0].click_action.target_slide==prs.slides[1]
    assert any(s.click_action.action==PP_ACTION.LAST_SLIDE_VIEWED for s in prs.slides[1].shapes)
    actions={s.click_action.action for s in prs.slides[2].shapes}
    assert {PP_ACTION.HYPERLINK,PP_ACTION.NEXT_SLIDE,PP_ACTION.PREVIOUS_SLIDE}<=actions
    assert any(s.name.startswith('!!details') for s in linked)
    html,_=render_document(deck,assets=services.assets_lookup(),inline=True)
    assert 'data-action=' in html and '#s-'+detail['id'] in html
    assert '.slide-action>.b{left:0!important;top:0!important;' in html
    document=present_document(html)
    assert 'routeHistory.push(origin)' in document and 'const origin=deck.getIndices()' in document
    assert 'event.stopImmediatePropagation()' in document and 'deck.toggleOverview(false)' in document

def test_routes_and_live_html_use_authoritative_navigation(client):
    d,first,detail=route_deck(client.svc)
    base=f"/api/decks/{d['id']}"
    result=client.put(base+'/navigation',json={'action':'set','slide_id':first['id'],'block_index':0,
        'target':{'kind':'slide','slide_id':detail['id']}})
    assert result.status_code==200 and client.get(base+'/navigation').json()['links'][1]['valid']
    room=client.post(base+'/live',json={}).json()
    page=client.get(room['audience_path']).text
    assert 'data-session-control=' in page and room['session_id']+'/control' in page
    assert 'PRIVATE' not in page

def test_template_native_actions_keep_destination_and_branding(services,tmp_path):
    from test_templates import template
    file=tmp_path/'brand.pptx';binding=template(file)
    d,first,detail=route_deck(services)
    navigation.run(services,navigation.NavigationArgs(deck_id=d['id'],action='set',slide_id=first['id'],block_index=0,target=NavigationAction(kind='slide',slide_id=detail['id'])))
    call_tool(services,'deck_template',{'deck_id':d['id'],'path':str(file),'bindings':{'bullets':binding}})
    prs=Presentation(io.BytesIO(build_pptx(decks.deck_view(services,d['id']),assets=services.assets_lookup())))
    assert any(s.name=='Empty decorative banner' for s in prs.slides[0].shapes)
    link=next(s for s in prs.slides[0].shapes if s.click_action.action==PP_ACTION.NAMED_SLIDE)
    assert link.click_action.target_slide==prs.slides[1]

def test_navigation_session_undo_restores_link_and_inspection_is_a_noop(client):
    from test_agent_accountability import agent_call,undo
    d,first,detail=route_deck(client.svc)
    args={'deck_id':d['id'],'action':'set','slide_id':first['id'],'block_index':0,
          'target':{'kind':'slide','slide_id':detail['id']}}
    response=agent_call(client,'deck_navigation',args,session='route-edit')
    assert response.status_code==200
    reverted=undo(client,'route-edit',confirm=True).json()
    assert reverted['complete'],reverted
    assert 'action' not in decks.get_slide(client.svc,d['id'],first['id'])['blocks'][0]
    assert agent_call(client,'deck_navigation',{'deck_id':d['id']},session='route-read').status_code==200
    assert undo(client,'route-read',confirm=True).json()['complete']

def test_navigation_undo_does_not_overwrite_a_later_manual_edit(client):
    from test_agent_accountability import agent_call,undo
    d,first,detail=route_deck(client.svc)
    assert agent_call(client,'deck_navigation',{'deck_id':d['id'],'action':'set','slide_id':first['id'],'block_index':0,
        'target':{'kind':'slide','slide_id':detail['id']}},session='route-edit').status_code==200
    decks.patch_slide(client.svc,d['id'],first['id'],{'notes':'Later manual notes'})
    assert not undo(client,'route-edit',confirm=True).json()['complete']
    assert decks.get_slide(client.svc,d['id'],first['id'])['notes']=='Later manual notes'
