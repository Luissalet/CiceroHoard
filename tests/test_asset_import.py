from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import pytest
from pptx import Presentation
from cicero_hoard import decks, agent_tools
from cicero_hoard.errors import CiceroError, Refused
from conftest import png_bytes, new_deck


def test_local_import_retry_and_native_export(services, tmp_path):
    deck = new_deck(services)
    file = tmp_path/'figure.png'; file.write_bytes(png_bytes())
    first = agent_tools.call_tool(services, 'asset_import', {'deck_id':deck['id'],'path':str(file)})
    with ThreadPoolExecutor(max_workers=4) as pool:
        repeated = list(pool.map(lambda _:agent_tools.call_tool(services,'asset_import',{'deck_id':deck['id'],'path':str(file)}), range(4)))
    assert all(x['asset_id']==first['asset_id'] and x['reused'] for x in repeated)
    assert len(services.db.query('SELECT id FROM assets'))==1
    decks.add_slide(services,deck['id'],layout='image_text',title='Existing figure',blocks=[{'type':'image','asset_id':first['asset_id']}],notes='Source notes')
    exported=agent_tools.call_tool(services,'deck_export',{'deck_id':deck['id'],'format':'pptx'})
    assert any(s.shape_type==13 for s in Presentation(exported['path']).slides[0].shapes)
    assert file.read_bytes()==png_bytes()
    other=new_deck(services)
    second=agent_tools.call_tool(services,'asset_import',{'deck_id':other['id'],'path':str(file)})
    assert second['asset_id']!=first['asset_id']


def test_import_failure_and_scope_do_not_add_assets(services,tmp_path):
    deck=new_deck(services); file=tmp_path/'bad.png';file.write_bytes(b'not an image')
    with pytest.raises(CiceroError):agent_tools.call_tool(services,'asset_import',{'deck_id':deck['id'],'path':str(file)})
    services.config.file_roots=(tmp_path/'allowed',)
    with pytest.raises(Refused):agent_tools.call_tool(services,'asset_import',{'deck_id':deck['id'],'path':str(file)})
    assert not services.db.query('SELECT id FROM assets')


def test_import_route_matches_ui_upload(client,tmp_path):
    deck=client.post('/api/decks',json={'title':'Reusable import'}).json()
    file=tmp_path/'figure.png';file.write_bytes(png_bytes())
    imported=client.post(f"/api/decks/{deck['id']}/assets/import",json={'path':str(file)})
    assert imported.status_code==201
    uploaded=client.post(f"/api/decks/{deck['id']}/assets",files={'file':('renamed.png',file.read_bytes(),'image/png')})
    assert uploaded.status_code==201 and uploaded.json()['asset_id']==imported.json()['asset_id']
